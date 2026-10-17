"""Pitch v8: conserva v7 y añade QR, acceso Auditor y seguridad en diez líneas.

Requiere las dependencias de generar_v7.py, qrcode y python-docx; soffice en PATH.
La credencial se lee de ~/.config/istpetdev/private/jurado-pitch-v8.json.
Los PPTX/PDF/render-v8 contienen la credencial destinada al jurado: no versionar.
El guion versionable vive en ../guion-pitch-v8.md y no contiene la contraseña.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import qrcode
from docx import Document
from docx.shared import Pt

import generar_v7 as v7
from generar_v3 import (
    ACERO, AZUL, BLANCO, DORADO, DORADO_CLARO, GRIS_NIEBLA, MARGIN,
    PP_ALIGN, Timeline, picture, record, rect, text,
)

HERE = Path(__file__).resolve().parent
SCRIPT = HERE.parent / "guion-pitch-v8.md"
URL = "https://istpetdev.doicela.dev/"
CREDENTIAL = Path.home() / ".config/istpetdev/private/jurado-pitch-v8.json"


def security_lines():
    section = SCRIPT.read_text().split("## Seguridad en diez líneas", 1)[1].split("\n## ", 1)[0]
    lines = re.findall(r"^\d+\. (.+)$", section, re.M)
    assert len(lines) == 10, "La explicación de seguridad debe tener diez líneas"
    return lines


def slide_access(deck, credential, qr):
    slide = deck.slide(dark=True)
    sh = slide.shapes
    v7.chrome(deck, slide, True, "Acceso para el jurado", "Pruébenlo ustedes",
              "Cuenta del jurado · Auditor · Solo lectura · Datos sintéticos", 0)
    picture(sh, qr, MARGIN, 1.90, w=4.15, h=4.15).click_action.hyperlink.address = URL
    text(sh, MARGIN, 6.13, 4.4, 0.4, "istpetdev.doicela.dev", 21,
                DORADO_CLARO, bold=True, align=PP_ALIGN.CENTER)
    x, w = 5.40, 7.15
    rect(sh, x, 1.90, w, 1.65, ACERO)
    text(sh, x + .22, 2.01, w - .44, .23, "USUARIO DEL JURADO", 11, DORADO, bold=True)
    text(sh, x + .22, 2.29, w - .44, .41, credential["email"], 23, BLANCO, bold=True)
    text(sh, x + .22, 2.80, 1.25, .24, "Contraseña", 11, DORADO, bold=True)
    text(sh, x + 1.55, 2.75, w - 1.77, .42, credential["password"], 22, BLANCO,
         bold=True, label="credencial del jurado")
    text(sh, x + .22, 3.20, w - .44, .25,
         "Torre de Control → usar esta cuenta", 12, GRIS_NIEBLA)
    text(sh, x, 3.73, w, .35, "SEGURIDAD, EN PALABRAS SENCILLAS", 16, DORADO_CLARO, bold=True)
    for index, line in enumerate(security_lines()):
        text(sh, x, 4.14 + index * .235, w, .24, f"{index + 1:02d}  {line}", 12.6,
             BLANCO, label=f"seguridad {index + 1}")
    # La última lámina queda visible de inmediato, sin clics que oculten el QR.
    record(deck, slide, Timeline(), "Acceso del jurado · QR y seguridad", "Sin clics; dejar visible durante preguntas.")
    return slide


def build(animated, output, credential, qr, export=False, render_dir=None):
    deck = v7.build(animated, output, export=False)
    slide_access(deck, credential, qr)
    deck.prs.core_properties.title = "IstpetDev · Pitch Expo Clean · v8"
    deck.prs.core_properties.subject = "Caso logístico y acceso Auditor del jurado; guion v8"
    rows = script_rows()
    for slide, row in zip(deck.prs.slides, rows, strict=True):
        slide.notes_slide.notes_text_frame.text = f"{row[1]} · {row[2]} s\n{row[3]}"
    deck.prs.save(output)
    assert len(v7.v3.Presentation(str(output)).slides) == len(deck.notes) == 10
    if export:
        v7.v3.PDF_PATH = output.with_suffix(".pdf")
        v7.v3.RENDER_DIR = render_dir or HERE / "render-v8"
        v7.v3.export(output)
    return deck


def script_rows():
    rows = []
    for line in SCRIPT.read_text().splitlines():
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) == 5 and cells[0].isdigit():
            rows.append((int(cells[0]), cells[2], int(cells[3]), cells[4].strip("«»")))
    assert [r[0] for r in rows] == list(range(1, 11))
    assert sum(r[2] for r in rows) == 290
    return rows


def export_script():
    doc = Document()
    doc.styles["Normal"].font.name = "Open Sans"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Guion del pitch v8 — IstpetDev", 0)
    doc.add_paragraph("10 diapositivas · Meta: 4:50 · Tiempos propuestos, pendientes de ensayo.")
    elapsed = 0
    for number, voice, seconds, spoken in script_rows():
        end = elapsed + seconds
        stamp = lambda t: f"{t // 60:02d}:{t % 60:02d}"
        doc.add_heading(f"{stamp(elapsed)}–{stamp(end)} · {voice} · Diapositiva {number}", 2)
        doc.add_paragraph(spoken)
        elapsed = end
    doc.add_page_break()
    doc.add_heading("Seguridad en diez líneas — apoyo para preguntas", 1)
    for i, line in enumerate(security_lines(), 1):
        doc.add_paragraph(f"{i}. {line}")
    doc.add_paragraph("Fuente: documentación oficial de Amazon Cognito; controles de la aplicación verificados por separado.")
    doc.add_paragraph("https://docs.aws.amazon.com/cognito/latest/developerguide/cognito-how-to-authenticate.html")
    doc.add_paragraph("La credencial está únicamente en la última diapositiva del pitch entregado al jurado.")
    output = HERE / "guion-pitch-v8.docx"
    doc.save(output)
    with tempfile.TemporaryDirectory(prefix="guion-v8-") as tmp:
        subprocess.run([shutil.which("soffice"), f"-env:UserInstallation={Path(tmp).as_uri()}/lo",
                        "--headless", "--convert-to", "pdf", "--outdir", str(HERE), str(output)],
                       check=True, capture_output=True)
    assert output.with_suffix(".pdf").exists()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sin-credenciales", action="store_true", help="Generar copia de revisión versionable")
    args = parser.parse_args()
    os.umask(0o077)
    credential = ({"email": "jurado@istpetdev.doicela.dev", "password": "Entregada al jurado", "url": URL}
                  if args.sin_credenciales else json.loads(CREDENTIAL.read_text()))
    assert credential["email"] == "jurado@istpetdev.doicela.dev"
    assert credential["url"] == URL and credential["password"]
    stem = "pitch-v8-revision" if args.sin_credenciales else "pitch-v8"
    render_dir = HERE / ("render-v8-revision" if args.sin_credenciales else "render-v8")
    with tempfile.TemporaryDirectory(prefix="pitch-v8-") as tmp:
        qr = Path(tmp) / "qr.png"
        code = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_H, box_size=20, border=4)
        code.add_data(URL)
        code.make(fit=True)
        code.make_image(fill_color="black", back_color="white").save(qr)
        static = HERE / f"{stem}-estatico.pptx"
        build(False, static, credential, qr, export=True, render_dir=render_dir)
        static.unlink()
        static.with_suffix(".pdf").replace(HERE / f"{stem}.pdf")
        deck = build(True, HERE / f"{stem}.pptx", credential, qr)
        for file in (HERE / f"{stem}.pdf", HERE / f"{stem}.pptx"):
            file.chmod(0o600)
    export_script()
    print(f"Pitch v8: {len(deck.notes)} diapositivas; guion 290 s. Credencial omitida del log.")
    for warning in v7.v3.FIT_WARNINGS:
        print(f"AVISO de ajuste · {warning.replace(credential['password'], '[credencial omitida]')}")


if __name__ == "__main__":
    main()
