"""Genera el entregable 2 (Solución y propuesta de valor) en Word.

Uso: python3 generar.py  (desde esta carpeta). Luego convertir a PDF con LibreOffice.
Las capturas se leen de img/captura-1..4.png; reemplazar esos archivos y volver a generar.
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
IMG = AQUI / "img"
CYAN, LIME, ORANGE, DARK, GREY = "00A7C4", "9EC200", "EA601F", "2D3748", "64748B"
MIST = "F8FAFC"


def rgb(h):
    return RGBColor.from_string(h)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def borders(table, color="CBD5E1", size="4"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single" if size != "0" else "nil")
        e.set(qn("w:sz"), size)
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def left_bar(paragraph, color, size="36"):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    bdr.append(left)
    pPr.append(bdr)


def bottom_rule(paragraph, color=CYAN, size="12"):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), size)
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), color)
    bdr.append(b)
    pPr.append(bdr)


def fix_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips)))
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
        for c, w in zip(row.cells, widths):
            c.width = w


def vcenter(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), "center")
    tcPr.append(v)


def para(doc_or_cell, text="", size=10.5, bold=False, color=DARK, align=None, after=4, italic=False):
    p = doc_or_cell.add_paragraph()
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = rgb(color)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    return p


def rich(p, parts, size=10.5):
    """parts: lista de (texto, negrita)."""
    for text, b in parts:
        r = p.add_run(text)
        r.bold = b
        r.font.size = Pt(size)
        r.font.color.rgb = rgb(DARK)
    return p


def h1(doc, num, text):
    p = doc.add_heading(level=1)
    r = p.add_run(f"{num}  ")
    r.font.color.rgb = rgb(CYAN)
    r2 = p.add_run(text)
    r2.font.color.rgb = rgb(DARK)
    for run in (r, r2):
        run.font.size = Pt(17)
        run.bold = True
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    bottom_rule(p)
    return p


def bullet(doc, parts):
    p = doc.add_paragraph(style="List Bullet")
    rich(p, parts)
    p.paragraph_format.space_after = Pt(3)
    return p


def chip(cell, text, color):
    cell.text = ""
    shade(cell, color)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = rgb("FFFFFF")


def footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("IstpetDev · Hackathon Expo Clean Ecuador 2026 · Reto 1 · Entregable 2 — Solución y propuesta de valor · pág. ")
    r.font.size = Pt(8)
    r.font.color.rgb = rgb(GREY)
    run = p.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), tag)
            run._r.append(f)
        else:
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = txt
            run._r.append(t)
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(GREY)


doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(10.5)
st.font.color.rgb = rgb(DARK)
for s in doc.sections:
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.top_margin = s.bottom_margin = Cm(2.0)
    s.left_margin = s.right_margin = Cm(2.2)

# ---------------- PORTADA ----------------
cover = doc.sections[0]
cover.different_first_page_header_footer = True
para(doc, "", after=40)
doc.add_picture(str(IMG / "logo.png"), width=Cm(11))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
para(doc, "", after=50)
para(doc, "HACKATHON EXPO CLEAN ECUADOR 2026", 11, True, CYAN, after=2)
para(doc, "Reto 1 — Optimización y trazabilidad logística del abastecimiento de insumos", 11, False, GREY, after=26)
t = para(doc, "Definición de la solución", 30, True, DARK, after=0)
para(doc, "y propuesta de valor", 30, True, CYAN, after=14)
p = para(doc, "Entregable 2 · Solución (Manual §9 y Anexos 5.2 y 5.6)", 12, False, DARK, after=60)
left_bar(p, ORANGE, "24")

tb = doc.add_table(rows=1, cols=2)
tb.alignment = WD_TABLE_ALIGNMENT.LEFT
borders(tb, "FFFFFF", "0")
c1, c2 = tb.rows[0].cells
c1.width, c2.width = Cm(8.3), Cm(8.3)
para(c1, "EQUIPO", 9, True, GREY, after=2)
para(c1, "IstpetDev", 13, True, DARK, after=8)
for name in ("Bryan Baño", "Jorge Doicela", "Joseph Baño", "Anthony Pilatasig", "Alexander Martínez"):
    para(c1, name, 10.5, False, DARK, after=1)
para(c2, "FECHA", 9, True, GREY, after=2)
para(c2, "16 de octubre de 2026", 11, False, DARK, after=8)
para(c2, "LUGAR", 9, True, GREY, after=2)
para(c2, "Quito, Ecuador", 11, False, DARK, after=8)
para(c2, "VERSIÓN", 9, True, GREY, after=2)
para(c2, "Borrador para el Checkpoint de avances", 11, False, DARK, after=8)
for c in (c1, c2):
    c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)

# ---------------- CUERPO ----------------
body = doc.add_section(WD_SECTION.NEW_PAGE)
footer(body)

h1(doc, "1", "El problema")
para(doc, "Una empresa de limpieza abastece de insumos a decenas de puntos de servicio: hospitales, oficinas y "
          "comercios. Hoy la reposición se programa con calendarios fijos, semanales o quincenales, y el faltante "
          "suele conocerse cuando el cliente llama. Cada urgencia rompe las rutas planificadas, encarece el "
          "transporte y deja sin evidencia de qué se entregó, a quién y cuándo.", after=8)
q = para(doc, "", after=2)
rich(q, [("«No es lo mismo limpiar un quirófano, una sala de neonatos o una habitación.» ", False),
         ("«Muchas veces mis clientes no hacen compras programadas.»", False)], 10.5)
for r in q.runs:
    r.italic = True
left_bar(q, CYAN)
para(doc, "Testimonios de la charla del sector, Expo Clean, 16-oct-2026.", 8.5, False, GREY, after=10)

h1(doc, "2", "¿Qué proponemos?")
p = para(doc, "", after=8)
rich(p, [("IstpetDev", True), (" es una plataforma que permite a una empresa de limpieza ", False),
         ("reabastecer sus puntos antes de que falte un insumo", True), (", decidir ", False),
         ("desde qué bodega y en qué vehículo", True), (" sale cada entrega, ", False),
         ("atender primero lo crítico", True), (" y ", False),
         ("demostrar la custodia de cada caja", True), (" desde la bodega hasta la recepción, incluso sin señal.", False)], 11)

h1(doc, "3", "¿Cómo funciona?")
doc.add_picture(str(IMG / "flujo.png"), width=Cm(16.6))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
para(doc, "Figura 1. Recorrido de un insumo en IstpetDev: anticipar, decidir y demostrar.", 8.5, False, GREY,
     WD_ALIGN_PARAGRAPH.CENTER, after=8)

pasos = [
    ("Anticipar", CYAN, [
        ("Datos", "Importa catálogo, puntos, bodegas, stock y consumo desde Excel o CSV, con validación por fila."),
        ("Riesgo", "Calcula los días de cobertura de cada punto y producto frente al tiempo de reposición, y explica por qué es Crítico, Alto o Vigilar."),
        ("Stock por bodega", "Busca en qué bodegas hay el producto y propone el origen, respetando stock de seguridad y vencimientos."),
    ]),
    ("Decidir", LIME, [
        ("Vehículo", "Elige el vehículo por capacidad de peso y volumen, bodega base y costo."),
        ("Ruta priorizada", "Atiende primero a los críticos y a los de menor holgura; respeta ventanas horarias y jornada, usa distancias por calles reales y suma puntos cercanos si el desvío es corto."),
        ("Agente y aprobación", "Un agente (n8n + DeepSeek) consulta y explica; el motor determinista calcula y una persona aprueba antes de publicar. Ante un cierre vial o un retraso, replanifica lo pendiente."),
    ]),
    ("Demostrar", ORANGE, [
        ("Custodia con QR", "Cada caja lleva un QR: despacho, entrega y recepción quedan con cantidad, lote, responsable, foto y firma."),
        ("Sin señal", "Conductor y receptor trabajan en una aplicación que funciona sin conexión y sincroniza sin duplicar."),
        ("Impacto", "Una simulación de 90 días compara la política con el calendario fijo y con la reacción tardía."),
    ]),
]
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
para(doc, "Paso a paso", 13, True, DARK, after=6)
tbl = doc.add_table(rows=0, cols=3)
borders(tbl, "E2E8F0", "4")
widths = (Cm(2.6), Cm(3.4), Cm(10.6))
for etapa, color, filas in pasos:
    for i, (nombre, texto) in enumerate(filas):
        row = tbl.add_row().cells
        for c, w in zip(row, widths):
            c.width = w
        if i == 0:
            chip(row[0], etapa.upper(), color)
        else:
            shade(row[0], color)
        row[1].text = ""
        r = row[1].paragraphs[0].add_run(nombre)
        r.bold = True
        r.font.size = Pt(10)
        row[2].text = ""
        r = row[2].paragraphs[0].add_run(texto)
        r.font.size = Pt(9.5)
    # unir la celda de etapa
    start = len(tbl.rows) - len(filas)
    tbl.cell(start, 0).merge(tbl.cell(len(tbl.rows) - 1, 0))
    m = tbl.cell(start, 0)
    for extra in m.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    vcenter(m)
fix_widths(tbl, widths)

h1(doc, "4", "¿Qué la hace diferente?")
difs = [
    (CYAN, "Cadena completa en una sola operación", "riesgo → bodega → vehículo → ruta → custodia → recepción. No es solo un ruteador ni solo un inventario."),
    (LIME, "Prioridad explicable", "cada decisión muestra los datos que la causan: cobertura, tiempo de reposición y criticidad del área. La IA explica; no inventa cantidades."),
    (ORANGE, "Funciona sin señal", "en el punto de entrega, donde la conectividad suele fallar."),
    (DARK, "Evidencia verificable", "cada caja se consulta escaneando su QR, con su historial de custodia."),
]
for color, titulo, texto in difs:
    p = para(doc, "", after=6)
    rich(p, [(titulo + ": ", True), (texto, False)])
    left_bar(p, color, "30")

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
h1(doc, "5", "El sistema en funcionamiento")
para(doc, "Capturas del prototipo desplegado (escenario demo-region).", 9.5, False, GREY, after=6)
caps = [
    ("captura-1.png", "Figura 2. Alertas de riesgo con la explicación del cálculo."),
    ("captura-2.png", "Figura 3. Ruta priorizada sobre el mapa, con paradas y carga."),
    ("captura-3.png", "Figura 4. Aplicación de campo: escaneo QR y recepción."),
    ("captura-4.png", "Figura 5. Simulación de 90 días frente a los métodos actuales."),
]
grid = doc.add_table(rows=2, cols=2)
borders(grid, "FFFFFF", "0")
for idx, (img, cap) in enumerate(caps):
    cell = grid.cell(idx // 2, idx % 2)
    cell.width = Cm(8.3)
    cell.paragraphs[0].add_run().add_picture(str(IMG / img), width=Cm(7.9))
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(cell, cap, 8.5, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, after=8)
fix_widths(grid, (Cm(8.3), Cm(8.3)))

h1(doc, "6", "Propuesta de valor")
box = doc.add_table(rows=1, cols=1)
borders(box, CYAN, "8")
cell = box.cell(0, 0)
shade(cell, "E0F7FB")
cell.text = ""
p = cell.paragraphs[0]
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(6)
rich(p, [("Para ", True), ("gerencias de operaciones de empresas de limpieza que abastecen muchos puntos, ", False),
         ("que hoy ", True), ("se enteran del faltante cuando el cliente llama y reponen con calendarios fijos sin poder comprobar qué se entregó, ", False),
         ("IstpetDev ", True), ("anticipa el faltante, decide desde dónde y en qué vehículo atender primero lo crítico, y demuestra cada entrega, ", False),
         ("a diferencia de ", True), ("hojas de cálculo, rutas fijas y listas de control en papel.", False)], 11)
para(doc, "", after=4)

ut = doc.add_table(rows=1, cols=3)
borders(ut)
hdr = ut.rows[0].cells
_h = ut.rows[0]._tr.get_or_add_trPr(); _x = OxmlElement("w:tblHeader"); _h.append(_x)
for c, txt, w in zip(hdr, ("Usuario", "Hoy", "Con IstpetDev"), (Cm(3.2), Cm(6.6), Cm(6.8))):
    c.width = w
    chip(c, txt, DARK)
for u, hoy, con in [
    ("Planificador", "Rutas fijas y urgencias por llamada", "Prioridad y ruta justificadas con datos"),
    ("Bodega", "Despachos sin conciliación", "Cantidad, lote y responsable por caja"),
    ("Conductor", "Sin señal ni registro de incidentes", "Trabajo sin conexión y replanificación"),
    ("Receptor", "Entrega difícil de verificar", "Recepción con QR, foto y firma"),
    ("Gerencia", "Sobrecostos sin explicación", "Indicadores con origen y comparación"),
]:
    row = ut.add_row().cells
    for c, txt, w, b in zip(row, (u, hoy, con), (Cm(3.2), Cm(6.6), Cm(6.8)), (True, False, False)):
        c.width = w
        c.text = ""
        r = c.paragraphs[0].add_run(txt)
        r.bold = b
        r.font.size = Pt(9.5)
    shade(row[2], "F1F8D9")
fix_widths(ut, (Cm(3.2), Cm(6.6), Cm(6.8)))

h1(doc, "7", "Alcance actual y próximos pasos")
bullet(doc, [("En vivo: ", True), ("una región con unos 15 puntos y 2 vehículos. La escala de 80 puntos y 90 días se presenta como simulación con datos sintéticos.", False)])
bullet(doc, [("En construcción durante el evento: ", True), ("hora límite para puntos críticos, consulta de stock y vehículos por el agente, y comparación del costo de priorizar.", False)])
bullet(doc, [("Próximos pasos: ", True), ("zonas de reparto habituales, conversión de cajas a unidades, consumo por superficie (litros por m²) y calibración del costo con viajes reales.", False)])
bullet(doc, [("Validación: ", True), ("aún sin cifras de ahorro confirmadas por una empresa; los resultados de la simulación se muestran con su escenario y semilla.", False)])

out = AQUI / "IstpetDev-Entregable-2-Solucion.docx"
doc.save(out)
print(out)
