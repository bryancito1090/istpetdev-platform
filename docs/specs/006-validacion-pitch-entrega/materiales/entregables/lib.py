"""Estilo común de los entregables oficiales (mismo formato que el entregable 2).

Lo usan e1_problema.py, e2 se conserva en ../entregable-2, e3_prototipo.py, e4_validacion.py,
e5_viabilidad.py y e6_modelo_negocio.py. El pitch (entregable 7) es pitch-v4.pdf.
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
IMG = AQUI.parent / "entregable-2" / "img"
CYAN, LIME, ORANGE, DARK, GREY = "00A7C4", "9EC200", "EA601F", "2D3748", "64748B"
ANCHO = Cm(16.6)
EQUIPO = ("Bryan Baño", "Jorge Doicela", "Joseph Baño", "Anthony Pilatasig", "Alexander Martínez")


def rgb(h):
    return RGBColor.from_string(h)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def borders(table, color="CBD5E1", size="4"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single" if size != "0" else "nil")
        e.set(qn("w:sz"), size)
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def _pbdr(paragraph, edge, color, size, space):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), size)
    e.set(qn("w:space"), space)
    e.set(qn("w:color"), color)
    bdr.append(e)
    pPr.append(bdr)


def left_bar(paragraph, color, size="36"):
    _pbdr(paragraph, "left", color, size, "8")


def bottom_rule(paragraph, color=CYAN, size="12"):
    _pbdr(paragraph, "bottom", color, size, "4")


def fix_widths(table, widths):
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips)))
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c, w in zip(row.cells, widths):
            c.width = w


def para(parent, text="", size=10.5, bold=False, color=DARK, align=None, after=4, italic=False):
    p = parent.add_paragraph()
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = rgb(color)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    return p


def rich(p, parts, size=10.5):
    """parts: lista de (texto, negrita)."""
    for text, b in parts:
        r = p.add_run(text)
        r.bold = b
        r.font.size = Pt(size)
        r.font.color.rgb = rgb(DARK)
    return p


def h1(doc, num, text):
    p = doc.add_heading(level=1)
    r = p.add_run(f"{num}  ")
    r.font.color.rgb = rgb(CYAN)
    r2 = p.add_run(text)
    r2.font.color.rgb = rgb(DARK)
    for run in (r, r2):
        run.font.size = Pt(17)
        run.bold = True
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    bottom_rule(p)
    return p


def h2(doc, text):
    p = para(doc, text, 12.5, True, DARK, after=4)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, parts):
    p = doc.add_paragraph(style="List Bullet")
    rich(p, parts if isinstance(parts, list) else [(parts, False)])
    p.paragraph_format.space_after = Pt(3)
    return p


def aviso(doc, titulo, texto, color=ORANGE):
    """Recuadro con barra lateral para límites, supuestos y pendientes."""
    p = para(doc, "", after=8)
    rich(p, [(titulo + " ", True), (texto, False)], 10)
    left_bar(p, color, "30")
    return p


def cita(doc, texto, fuente):
    q = para(doc, "", after=2)
    r = q.add_run(texto)
    r.italic = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = rgb(DARK)
    left_bar(q, CYAN)
    para(doc, fuente, 8.5, False, GREY, after=8)


def chip(cell, text, color, size=8.5):
    cell.text = ""
    shade(cell, color)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = rgb("FFFFFF")


def tabla(doc, encabezados, filas, anchos, negrita_col0=True, tam=9.5, zebra=False):
    t = doc.add_table(rows=1, cols=len(encabezados))
    borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0]
    hdr._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for c, txt in zip(hdr.cells, encabezados):
        chip(c, txt, DARK, 9)
    for n, fila in enumerate(filas):
        cells = t.add_row().cells
        for i, (c, txt) in enumerate(zip(cells, fila)):
            c.text = ""
            r = c.paragraphs[0].add_run(str(txt))
            r.font.size = Pt(tam)
            r.bold = negrita_col0 and i == 0
            r.font.color.rgb = rgb(DARK)
            if zebra and n % 2:
                shade(c, "F8FAFC")
    fix_widths(t, anchos)
    para(doc, "", after=4)
    return t


def figura(doc, ruta, ancho, pie):
    doc.add_picture(str(ruta), width=ancho)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, pie, 8.5, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, after=8)


def salto(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _pie(section, texto):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"IstpetDev · Hackathon Expo Clean Ecuador 2026 · Reto 1 · {texto} · pág. ")
    r.font.size = Pt(8)
    r.font.color.rgb = rgb(GREY)
    run = p.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), tag)
            run._r.append(f)
        else:
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = txt
            run._r.append(t)
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(GREY)


def nuevo(numero, nombre, titulo1, titulo2, referencia, version="Versión para la entrega del 17 de octubre de 2026"):
    """Crea el documento con la portada del formato establecido y deja el cursor en el cuerpo."""
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(10.5)
    st.font.color.rgb = rgb(DARK)
    for s in doc.sections:
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.top_margin = s.bottom_margin = Cm(2.0)
        s.left_margin = s.right_margin = Cm(2.2)
    doc.sections[0].different_first_page_header_footer = True
    para(doc, "", after=40)
    doc.add_picture(str(IMG / "logo.png"), width=Cm(11))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
    para(doc, "", after=50)
    para(doc, "HACKATHON EXPO CLEAN ECUADOR 2026", 11, True, CYAN, after=2)
    para(doc, "Reto 1 — Optimización y trazabilidad logística del abastecimiento de insumos", 11, False, GREY, after=26)
    para(doc, titulo1, 30, True, DARK, after=0)
    para(doc, titulo2, 30, True, CYAN, after=14)
    p = para(doc, f"Entregable {numero} · {nombre} ({referencia})", 12, False, DARK, after=60)
    left_bar(p, ORANGE, "24")
    tb = doc.add_table(rows=1, cols=2)
    tb.alignment = WD_TABLE_ALIGNMENT.LEFT
    borders(tb, "FFFFFF", "0")
    c1, c2 = tb.rows[0].cells
    c1.width, c2.width = Cm(8.3), Cm(8.3)
    para(c1, "EQUIPO", 9, True, GREY, after=2)
    para(c1, "IstpetDev", 13, True, DARK, after=8)
    for n in EQUIPO:
        para(c1, n, 10.5, False, DARK, after=1)
    for rot, val in (("FECHA", "17 de octubre de 2026"), ("LUGAR", "Quito, Ecuador"), ("VERSIÓN", version)):
        para(c2, rot, 9, True, GREY, after=2)
        para(c2, val, 11, False, DARK, after=8)
    for c in (c1, c2):
        c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    _pie(body, f"Entregable {numero} — {nombre}")
    return doc


def guardar(doc, nombre):
    out = AQUI / nombre
    doc.save(out)
    print(out)
