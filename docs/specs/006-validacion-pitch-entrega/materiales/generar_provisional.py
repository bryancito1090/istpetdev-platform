"""Genera la presentación provisional v1 de siete bloques del pitch (spec 006).

Solo usa frases autorizadas por matriz-afirmaciones.md (usables con su límite
o frase de sustitución) y la «Frase vigente» de escena-y-contingencias.md.
No es formato oficial: O-04/O-06 siguen pendientes.

Uso:
    uv run --with python-pptx python docs/specs/006-validacion-pitch-entrega/materiales/generar_provisional.py
    soffice --headless --convert-to pdf docs/specs/006-validacion-pitch-entrega/materiales/pitch-provisional-v1.pptx --outdir docs/specs/006-validacion-pitch-entrega/materiales/
"""

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

VERSION = "v1 · 16-oct-2026"
AVISO = (f"PROVISIONAL {VERSION} · IstpetDev · no es formato oficial "
         "(O-04/O-06 pendientes) · no entregado")

TINTA = RGBColor(0x1F, 0x29, 0x37)
GRIS = RGBColor(0x6B, 0x72, 0x80)
ROTULO = RGBColor(0xB4, 0x53, 0x09)
SLATE_500 = RGBColor(0x64, 0x74, 0x8B)
SLATE_700 = RGBColor(0x33, 0x41, 0x55)

DIAPOSITIVAS = [
    {
        "bloque": "1 · Problema",
        "titulo": "Reponer antes de la llamada",
        "rotulo": "Caso sintético · SERCOP",
        "lineas": [
            "En contratos públicos de limpieza hospitalaria, el proveedor está "
            "obligado a mantener los químicos de desinfección y los insumos "
            "fungibles con frecuencias de reposición definidas por el hospital.",
            "Nuestra hipótesis es que hoy la reposición empieza cuando el "
            "hospital ya avisó la falta.",
            "Caso sintético: una caja de desinfectante que va a un hospital.",
        ],
        "fuente": "Fuente: fichas técnicas SERCOP (bóveda, 99 Fuentes). Matriz A-01, A-02.",
    },
    {
        "bloque": "2 · Solución",
        "titulo": "Calcula, propone, una persona decide",
        "rotulo": "Propuesta · flujo entregable 2",
        "es_diagrama": True,
        "lineas": [
            "El motor calcula alertas explicables y propone rutas; una persona aprueba antes de cualquier despacho.",
            "Cadena integrada de 8 etapas (Entregable 2) con replanificación ante incidentes y custodia física con código QR.",
        ],
        "fuente": "Entregable 2 (Solución y propuesta de valor). Matriz A-06, A-10 (sustitución), A-11 (sustitución).",
    },
    {
        "bloque": "3 · Prototipo / demo",
        "titulo": "Una caja hacia el hospital norte",
        "rotulo": "Pruebas automatizadas · hora límite en construcción (#193/#196)",
        "lineas": [
            "Escenario sintético: una bodega, 15 puntos y dos vehículos.",
            "En nuestras pruebas, el hospital norte tiene cinco días de cobertura "
            "de desinfectante: riesgo alto.",
            "La propuesta asigna una ruta a cada vehículo respetando capacidad, "
            "horarios de atención y jornada (hora límite para críticos en "
            "construcción #193/#196).",
            "Ante un cierre de vía se conservan las paradas hechas y se crea una "
            "nueva versión de la ruta.",
            "Diseño: la caja queda ligada a su lote y a la entrega.",
            "La consulta pública devuelve producto, estado e hitos sin datos "
            "privados; un token alterado recibe 404 en la prueba de API.",
        ],
        "fuente": "Matriz A-03 (sust.), A-04, A-07, A-09, A-13 (propuesta), A-16 (sust.). Issues #193, #196.",
    },
    {
        "bloque": "4 · Impacto",
        "titulo": "Cómo mediremos el impacto",
        "rotulo": "Simulación 90 días · resultado pendiente",
        "lineas": [
            "Diseñamos la comparación antes de ver resultados: 90 días "
            "sintéticos, tres semillas y un escenario adverso aparte.",
            "Contra reposición reactiva y calendario fijo.",
            "Contamos los quiebres evitados y también los nuevos, sin restar "
            "unos de otros.",
            "El resultado solo lo mostraremos con su archivo verificable.",
        ],
        "fuente": "Matriz A-19 (sustitución), A-20. A-21 retirada: sin corrida oficial de 005.",
    },
    {
        "bloque": "5 · Viabilidad",
        "titulo": "Qué ya funciona y qué no",
        "rotulo": "Prueba local",
        "lineas": [
            "Cargamos archivos de otro formato escribiendo solo un perfil, "
            "sin programar.",
            "En prueba local, la ruta de 80 puntos se calcula en menos de "
            "cinco segundos.",
            "Límite: es una heurística; no prometemos la ruta óptima.",
        ],
        "fuente": "Matriz A-22, A-24, A-33.",
    },
    {
        "bloque": "6 · Modelo de negocio",
        "titulo": "Quién paga y por qué",
        "rotulo": "Hipótesis",
        "lineas": [
            "Nuestro cliente por validar es la empresa de limpieza.",
            "Cobraríamos implementación y una suscripción.",
            "El precio lo definiremos con costos reales del piloto.",
            "El costo de operación lo mediremos en el piloto.",
        ],
        "fuente": "Matriz A-31 (propuesta), A-30 y A-26 (sustitución).",
    },
    {
        "bloque": "7 · Próximo paso",
        "titulo": "Piloto de cuatro semanas",
        "rotulo": "Propuesta",
        "lineas": [
            "Semana 1: acordar bodega, puntos y productos; medir la operación "
            "actual sin cambiarla.",
            "Semana 2: recomendaciones en paralelo, con aprobación humana.",
            "Semana 3: operar un grupo pequeño con caja y QR.",
            "Semana 4: comparar contra la operación actual y decidir si seguir.",
            "Buscamos una empresa que quiera medirlo con nosotros.",
        ],
        "fuente": "Matriz A-32; ficha de piloto del plan 006.",
    },
]


def texto(slide, x, y, w, h, contenido, tam, color=TINTA, negrita=False, align=PP_ALIGN.LEFT, space_after=8):
    caja = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    marco = caja.text_frame
    marco.word_wrap = True
    parrafos = contenido if isinstance(contenido, list) else [contenido]
    for i, linea in enumerate(parrafos):
        p = marco.paragraphs[0] if i == 0 else marco.add_paragraph()
        p.text = linea
        p.font.size = Pt(tam)
        p.font.color.rgb = color
        p.font.bold = negrita
        p.alignment = align
        p.space_after = Pt(space_after)
    return caja


def crear_tarjeta(slide, x, y, w, h, paso, subtitulo, detalle, bg_color, border_color, paso_color, sub_color, border_w=1.2):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(border_w)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.10)
    tf.margin_right = Inches(0.10)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.06)

    p0 = tf.paragraphs[0]
    p0.text = paso
    p0.font.bold = True
    p0.font.size = Pt(11)
    p0.font.color.rgb = paso_color
    p0.space_after = Pt(1)

    p1 = tf.add_paragraph()
    p1.text = subtitulo
    p1.font.bold = True
    p1.font.size = Pt(10)
    p1.font.color.rgb = sub_color
    p1.space_after = Pt(2)

    p2 = tf.add_paragraph()
    p2.text = detalle
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = SLATE_700
    p2.space_after = Pt(0)
    return shape


def crear_flecha(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0.15)
    tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.text = "→"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = SLATE_500
    p.alignment = PP_ALIGN.CENTER


def crear_banner(slide, x, y, w, h, titulo, texto_b, bg, border, tit_col):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg
    shape.line.color.rgb = border
    shape.line.width = Pt(1.2)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.06)

    p0 = tf.paragraphs[0]
    p0.text = titulo
    p0.font.bold = True
    p0.font.size = Pt(10)
    p0.font.color.rgb = tit_col
    p0.space_after = Pt(2)

    p1 = tf.add_paragraph()
    p1.text = texto_b
    p1.font.size = Pt(8.5)
    p1.font.color.rgb = SLATE_700
    p1.space_after = Pt(0)
    return shape


def construir_diapositiva_solucion(slide, d):
    texto(slide, 0.6, 1.60, 12.1, 0.65, ["• " + l for l in d["lineas"]], 15, SLATE_700, False, space_after=4)

    # Fase 1
    texto(slide, 0.6, 2.35, 12.1, 0.25, "FASE 1: PREPARACIÓN DE DATOS Y CÁLCULO DE DISPONIBILIDAD", 10, SLATE_500, True)

    pasos_r1 = [
        ("1 · Datos", "Stock y consumo", "Catálogo, puntos y bodegas desde Excel/CSV validado."),
        ("2 · Riesgo", "Por punto y producto", "Días de cobertura vs lead time (Crítico / Alto / Vigilar)."),
        ("3 · Bodega", "Stock por bodega", "Asignación de origen respetando seguridad y FEFO."),
        ("4 · Vehículo", "Vehículo disponible", "Elección por capacidad peso/volumen, base y costo."),
    ]

    xs = [0.60, 3.70, 6.80, 9.90]
    card_w = 2.65
    card_h = 1.05
    y_r1 = 2.65

    for i, (paso, sub, det) in enumerate(pasos_r1):
        crear_tarjeta(slide, xs[i], y_r1, card_w, card_h, paso, sub, det,
                      RGBColor(0xF0, 0xF9, 0xFF), RGBColor(0x02, 0x84, 0xC7),
                      RGBColor(0x0C, 0x4A, 0x6E), RGBColor(0x03, 0x69, 0xA1))
        if i < 3:
            crear_flecha(slide, xs[i] + card_w, y_r1, xs[i+1] - (xs[i] + card_w), card_h)

    # Fase 2
    texto(slide, 0.6, 3.82, 12.1, 0.25, "FASE 2: OPTIMIZACIÓN DE RUTA, DECISIÓN HUMANA Y CUSTODIA EN CAMPO", 10, SLATE_500, True)

    pasos_r2 = [
        ("5 · Ruta priorizada", "Críticos primero", "Ventanas horarias, jornada y menor holgura.",
         RGBColor(0xFE, 0xF3, 0xC7), RGBColor(0xD9, 0x77, 0x06), RGBColor(0x78, 0x35, 0x0F), RGBColor(0xB4, 0x53, 0x09), 1.2),
        ("6 · Aprobación humana", "Una persona decide", "Motor calcula; IA explica; persona aprueba despacho.",
         RGBColor(0xFF, 0xED, 0xD5), RGBColor(0xEA, 0x58, 0x0C), RGBColor(0x7C, 0x2D, 0x12), RGBColor(0xC2, 0x41, 0x0C), 2.0),
        ("7 · Despacho con QR", "Custodia física", "Cada caja ligada a lote, vehículo y responsable.",
         RGBColor(0xEC, 0xFD, 0xF5), RGBColor(0x10, 0xB9, 0x81), RGBColor(0x06, 0x4E, 0x3B), RGBColor(0x04, 0x78, 0x57), 1.2),
        ("8 · Recepción", "Con o sin señal", "PWA offline en punto de entrega, firma y sin duplicar.",
         RGBColor(0xEC, 0xFD, 0xF5), RGBColor(0x10, 0xB9, 0x81), RGBColor(0x06, 0x4E, 0x3B), RGBColor(0x04, 0x78, 0x57), 1.2),
    ]

    y_r2 = 4.10
    for i, (paso, sub, det, bg, border, p_col, s_col, bw) in enumerate(pasos_r2):
        crear_tarjeta(slide, xs[i], y_r2, card_w, card_h, paso, sub, det, bg, border, p_col, s_col, bw)
        if i < 3:
            crear_flecha(slide, xs[i] + card_w, y_r2, xs[i+1] - (xs[i] + card_w), card_h)

    # Callouts inferiores
    y_callout = 5.30
    callout_h = 0.85
    crear_banner(slide, 0.60, y_callout, 5.80, callout_h,
                 "⚠️ Replanificación ante incidente vial o retraso",
                 "Conserva paradas ya hechas y recalcula las pendientes con prioridad a puntos críticos (reingresa a Paso 5).",
                 RGBColor(0xFF, 0xFB, 0xEB), RGBColor(0xF5, 0x9E, 0x0B), RGBColor(0x92, 0x40, 0x0E))

    crear_banner(slide, 6.75, y_callout, 5.80, callout_h,
                 "🔄 Custodia offline y cierre de ciclo de stock",
                 "Recepción física (con o sin señal) confirma custodia y sincroniza datos para actualizar stock (cierra ciclo en Paso 1).",
                 RGBColor(0xF0, 0xFD, 0xF4), RGBColor(0x22, 0xC5, 0x5E), RGBColor(0x06, 0x5F, 0x46))


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    vacia = prs.slide_layouts[6]
    for d in DIAPOSITIVAS:
        s = prs.slides.add_slide(vacia)
        # Encabezado
        texto(s, 0.6, 0.35, 5.5, 0.4, d["bloque"], 14, GRIS, True)
        if d.get("rotulo"):
            texto(s, 6.2, 0.35, 6.5, 0.4, d["rotulo"], 14, ROTULO, True, align=PP_ALIGN.RIGHT)
        texto(s, 0.6, 0.85, 12.1, 0.75, d["titulo"], 32, TINTA, True)

        if d.get("es_diagrama"):
            construir_diapositiva_solucion(s, d)
        else:
            tam = 18 if len(d["lineas"]) >= 6 else (20 if len(d["lineas"]) <= 4 else 19)
            space_after = 6 if len(d["lineas"]) >= 6 else 10
            texto(s, 0.6, 1.95, 12.1, 4.3, ["• " + l for l in d["lineas"]], tam, space_after=space_after)

        # Pie
        texto(s, 0.6, 6.35, 12.1, 0.35, d["fuente"], 11, GRIS)
        texto(s, 0.6, 6.75, 12.1, 0.35, AVISO, 11, ROTULO, True)

    output_path = Path(__file__).resolve().parent / "pitch-provisional-v1.pptx"
    prs.save(output_path)
    print(f"Presentación guardada exitosamente en: {output_path}")


if __name__ == "__main__":
    main()
