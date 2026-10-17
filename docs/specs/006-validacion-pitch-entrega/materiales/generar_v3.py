"""Genera pitch-v3.pptx, pitch-v3.pdf y render-v3/ con la identidad del ISTPET.

Mismo contenido y fuentes que v2 (guion T4 y matriz de afirmaciones); cambia el
diseño, se agrega portada y aparece «Chillo el gallo» como narrador en cinco
láminas. Las animaciones (transición y entradas) se inyectan como XML de
PresentationML porque python-pptx no las expone.

Uso:
    uv run --with python-pptx --with pillow --with numpy --with scipy \
        python docs/specs/006-validacion-pitch-entrega/materiales/generar_v3.py

Requiere `soffice` (LibreOffice) para el PDF, `pdftoppm` (poppler) para los PNG
y las fuentes Oswald y Open Sans instaladas (copia en assets/fuentes/, licencia OFL).
El script mide cada texto con esas fuentes y avisa si alguno no cabe en su caja.
"""
from __future__ import annotations

import math
import shutil
import subprocess
from pathlib import Path

import numpy as np
from lxml import etree
from PIL import Image, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from scipy import ndimage as ndi


HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
FONTS = ASSETS / "fuentes"
DERIVED = ASSETS / "derivados"
ROUTE_CAPTURE = HERE / "entregable-2" / "img" / "captura-rutas.png"
PPTX_PATH = HERE / "pitch-v3.pptx"
PDF_PATH = HERE / "pitch-v3.pdf"
RENDER_DIR = HERE / "render-v3"

W, H = 13.333, 7.5
MARGIN = 0.75
# Borde derecho del contenido cuando Chillo ocupa la franja derecha de la lámina.
CHILLO_RIGHT = 9.4

# Paleta institucional medida sobre los logos (azul #242D57, fondo #162451, dorado #C4A857).
NOCHE = RGBColor(0x16, 0x24, 0x51)         # fondos oscuros (igual al logo sobre azul)
AZUL = RGBColor(0x24, 0x2D, 0x57)          # azul institucional: títulos y texto
ACERO = RGBColor(0x2E, 0x3B, 0x6E)         # superficies sobre fondo oscuro
NIEBLA = RGBColor(0xED, 0xF0, 0xF7)        # superficies claras
LINEA = RGBColor(0xD5, 0xDB, 0xEA)         # bordes y divisores claros
DORADO = RGBColor(0xC4, 0xA8, 0x57)        # acento institucional (formas, cifras sobre azul)
DORADO_CLARO = RGBColor(0xE6, 0xD7, 0xA3)  # resaltes de texto sobre azul
DORADO_OSCURO = RGBColor(0x7D, 0x65, 0x20)  # texto dorado sobre blanco (5,6:1)
CREMA = RGBColor(0xF7, 0xF1, 0xDF)         # fila o caja resaltada
GRIS = RGBColor(0x5B, 0x65, 0x85)          # texto secundario sobre blanco (5,8:1)
GRIS_NIEBLA = RGBColor(0xB9, 0xC2, 0xDC)   # texto secundario sobre azul (8,4:1)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

TITLE_FONT = "Oswald"
BODY_FONT = "Open Sans"
EVENT_TAG = "HACKATHON EXPO CLEAN 2026 · RETO 1"


# ---------------------------------------------------------------------------
# Medición de texto con las fuentes reales
# ---------------------------------------------------------------------------

# Alto de línea de cada familia en em (ascendente + descendente de la fuente).
LINE_EM = {TITLE_FONT: 1.482, BODY_FONT: 1.362}
FIT_WARNINGS: list[str] = []
_FONT_CACHE: dict[tuple, ImageFont.FreeTypeFont] = {}


def _font(family, bold, italic):
    key = (family, bold, italic)
    if key not in _FONT_CACHE:
        if family == TITLE_FONT:
            path, name = FONTS / "Oswald.ttf", "Bold" if bold else "Regular"
        elif italic:
            path, name = FONTS / "OpenSans-Italic.ttf", "Bold Italic" if bold else "Italic"
        else:
            path, name = FONTS / "OpenSans.ttf", "Bold" if bold else "Regular"
        font = ImageFont.truetype(str(path), 1000)
        font.set_variation_by_name(name)
        _FONT_CACHE[key] = font
    return _FONT_CACHE[key]


def width_in(chunk, size, family=BODY_FONT, bold=False, italic=False):
    return _font(family, bold, italic).getlength(chunk) / 1000 * size / 72


def wrapped_lines(runs, box_w):
    """Líneas que ocupa un párrafo de tramos (texto, tamaño, familia, negrita, cursiva[, spc])."""
    words = []
    for run in runs:
        chunk, size, family, bold, italic = run[:5]
        spc = run[5] if len(run) > 5 else 0
        for i, word in enumerate(chunk.split(" ")):
            if i > 0:
                words.append(None)
            if word:
                words.append((word, size, family, bold, italic, spc))
    lines, cur, pending_space = 1, 0.0, 0.0
    for item in words:
        if item is None:
            pending_space = width_in(" ", 14)
            continue
        word, size, family, bold, italic, spc = item
        # spc es el espaciado entre letras en centésimas de punto.
        w = width_in(word, size, family, bold, italic) + len(word) * spc / 7200
        if cur > 0 and cur + pending_space + w > box_w:
            lines += 1
            cur = w
        else:
            cur += (pending_space if cur > 0 else 0) + w
        pending_space = 0.0
    return lines


def check_fit(label, paras, box_w, box_h, line_spacing=1.0):
    """Registra un aviso si el texto medido no cabe en la caja."""
    height = 0.0
    for runs in paras:
        size = max(r[1] for r in runs)
        family = runs[0][2]
        n = wrapped_lines(runs, box_w)
        height += n * size * LINE_EM[family] * line_spacing / 72
    if height > box_h + 0.03:
        text_value = " ".join(r[0] for p in paras for r in p)[:70]
        FIT_WARNINGS.append(f"{label}: «{text_value}» necesita {height:.2f} in y la caja tiene {box_h:.2f} in")
    return height


# ---------------------------------------------------------------------------
# Recursos gráficos derivados
# ---------------------------------------------------------------------------

def cut_out_chillo(src: Path) -> Image.Image:
    """Quita el fondo blanco de la mascota sin halo y conserva su sombra como transparencia."""
    rgb = np.asarray(Image.open(src).convert("RGB")).astype(np.float32)
    h, w, _ = rgb.shape
    low = rgb.min(2)
    sat = rgb.max(2) - low
    rows = np.arange(h)[:, None].repeat(w, 1)
    shadow_zone = rows >= int(h * 0.815)
    candidate = ((sat <= 30) & (low >= 200)) | (shadow_zone & (sat <= 34) & (low >= 45))
    labels, _ = ndi.label(candidate)
    edge_labels = set(np.unique(np.concatenate(
        [labels[0], labels[-1], labels[:, 0], labels[:, -1]]))) - {0}
    bg = np.isin(labels, list(edge_labels))
    # Huecos blancos encerrados entre brazos y cuerpo; se respeta el texto de la camiseta.
    white = (sat <= 22) & (low >= 228) & (rows >= int(h * 0.435))
    holes, count = ndi.label(white)
    sizes = ndi.sum(white, holes, range(1, count + 1))
    for idx, size in enumerate(sizes, 1):
        if size < 120:
            continue
        ys, xs = np.where(holes == idx)
        on_shirt = 0.435 * w <= xs.mean() <= 0.60 * w and 0.50 * h <= ys.mean() <= 0.61 * h
        if not on_shirt:
            bg |= holes == idx
    # Pela el reborde claro neutro que queda pegado al fondo.
    light = (sat <= 40) & (low >= 165)
    for _ in range(4):
        grow = ndi.binary_dilation(bg) & light & ~bg
        if not grow.any():
            break
        bg |= grow
    fg_mask = Image.fromarray((~bg * 255).astype(np.uint8))
    fg_mask = fg_mask.filter(ImageFilter.MinFilter(3)).filter(ImageFilter.GaussianBlur(0.9))
    a_fg = np.asarray(fg_mask).astype(np.float32) / 255
    lum = rgb.mean(2)
    a_shadow = np.where(bg & shadow_zone, np.clip((250 - lum) / 250, 0, 1) * 0.9, 0)
    a_shadow = np.where(a_shadow < 0.05, 0, a_shadow)
    alpha = np.maximum(a_fg, a_shadow)
    # Descontamina el borde mezclado con blanco para que no quede halo sobre azul.
    af = a_fg[..., None]
    unmixed = np.clip((rgb - (1 - af) * 255) / np.maximum(af, 1e-3), 0, 255)
    edge = (a_fg > 0.02) & (a_fg < 0.98)
    out = np.where(edge[..., None], unmixed, rgb)
    out = np.where((a_shadow > a_fg)[..., None], 0, out)
    img = Image.fromarray(np.dstack([out, alpha * 255]).astype(np.uint8), "RGBA")
    return img.crop(img.getbbox())


def clean_logo(src: Path, white: bool) -> Image.Image:
    """Quita el marco gris semitransparente del PNG original; `white` pasa el azul a blanco."""
    px = np.asarray(Image.open(src).convert("RGBA")).astype(np.int16).copy()
    px[:, :2, 3] = px[:, -2:, 3] = px[:2, :, 3] = px[-2:, :, 3] = 0
    if white:
        navy = px[..., 2] > px[..., 0] + 10
        px[navy, 0:3] = 255
    img = Image.fromarray(px.astype(np.uint8), "RGBA")
    return img.crop(img.getbbox())


def prepare_assets() -> dict[str, Path]:
    DERIVED.mkdir(exist_ok=True)
    paths = {
        "chillo": DERIVED / "chillo-cuerpo-entero.png",
        "chillo_half": DERIVED / "chillo-medio-cuerpo.png",
        "logo": DERIVED / "istpet-logo.png",
        "logo_white": DERIVED / "istpet-logo-blanco.png",
        "capture": DERIVED / "captura-rutas-recorte.png",
    }
    chillo = cut_out_chillo(ASSETS / "chillo-el-gallo.png")
    chillo.save(paths["chillo"], optimize=True)
    # Medio cuerpo: de la cresta a la cintura, para asomarse desde el borde inferior.
    chillo.crop((0, 0, chillo.width, int(chillo.height * 0.64))).save(
        paths["chillo_half"], optimize=True)
    clean_logo(ASSETS / "ISTPET-ORIGINAL.png", white=False).save(paths["logo"], optimize=True)
    clean_logo(ASSETS / "ISTPET-ORIGINAL.png", white=True).save(paths["logo_white"], optimize=True)
    # Mismo recorte que v2: conserva el mapa y deja fuera navegación, métricas y fecha.
    cap = Image.open(ROUTE_CAPTURE)
    cw, ch = cap.size
    cap.crop((int(cw * 0.415), int(ch * 0.397), int(cw * (1 - 0.207)),
              int(ch * (1 - 0.308)))).save(paths["capture"], optimize=True)
    return paths


# ---------------------------------------------------------------------------
# Primitivas de dibujo
# ---------------------------------------------------------------------------

def _plain(shape, fill, line, line_width):
    """Relleno y borde explícitos, sin el estilo del tema (que agrega sombra en Impress)."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
    return shape


def rect(shapes, x, y, w, h, fill, line=None, radius=None, line_width=0.75):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shape.adjustments[0] = radius
    return _plain(shape, fill, line, line_width)


def oval(shapes, x, y, d, fill, line=None, line_width=1.25):
    shape = shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    return _plain(shape, fill, line, line_width)


def _style_run(run, size, color, bold, font, italic=False, spacing=None):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    if spacing is not None:
        run._r.get_or_add_rPr().set("spc", str(spacing))


def text(shapes, x, y, w, h, value, size=18, color=AZUL, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=BODY_FONT,
         italic=False, spacing=None, line_spacing=None, space_after=0, label=""):
    """Caja de texto. `value` es un str, una lista de párrafos o, por párrafo,
    una lista de tramos (texto, {color, bold, font, size}) para resaltar palabras."""
    box = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    paras = value if isinstance(value, list) else [value]
    measured = []
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        measured_runs = []
        for chunk, over in runs:
            run = p.add_run()
            run.text = chunk
            r_size, r_bold = over.get("size", size), over.get("bold", bold)
            r_font, r_italic = over.get("font", font), over.get("italic", italic)
            _style_run(run, r_size, over.get("color", color), r_bold, r_font, r_italic, spacing)
            measured_runs.append((chunk, r_size, r_font, r_bold, r_italic, spacing or 0))
        measured.append(measured_runs)
    check_fit(label or "texto", measured, w, h, line_spacing or 1.0)
    return box


def picture(shapes, path, x, y, w=None, h=None):
    kwargs = {}
    if w is not None:
        kwargs["width"] = Inches(w)
    if h is not None:
        kwargs["height"] = Inches(h)
    return shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)


def soft_shadow(shape, blur=10, dist=3, alpha=22):
    """Sombra exterior suave (a:outerShdw) que respetan PowerPoint e Impress."""
    sp_pr = shape._element.spPr
    for old in sp_pr.findall(qn("a:effectLst")):
        sp_pr.remove(old)
    effect = etree.SubElement(sp_pr, qn("a:effectLst"))
    shdw = etree.SubElement(effect, qn("a:outerShdw"), blurRad=str(Pt(blur)),
                            dist=str(Pt(dist)), dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(shdw, qn("a:srgbClr"), val="162451")
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


# ---------------------------------------------------------------------------
# Animaciones (PresentationML)
# ---------------------------------------------------------------------------

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
EFFECTS = {"fade": (10, 0), "wipe": (22, 8), "zoom": (53, 16)}


class Timeline:
    """Secuencia de entradas de una lámina.

    mode="click" abre un nuevo clic; "after" se encadena tras el efecto anterior
    (también al inicio de la lámina, sin clic); "with" arranca junto al anterior.
    """

    def __init__(self):
        self.groups: list[dict] = []
        self.text_targets: set[int] = set()

    def add(self, shape, effect="fade", mode="after", dur=450, delay=0):
        item = {"spid": shape.shape_id, "effect": effect, "dur": dur, "delay": delay}
        if shape._element.tag == qn("p:sp") and shape.has_text_frame:
            self.text_targets.add(shape.shape_id)
        if mode == "click" or not self.groups:
            self.groups.append({"auto": mode != "click", "steps": [[item]]})
        elif mode == "with":
            self.groups[-1]["steps"][-1].append(item)
        else:
            self.groups[-1]["steps"].append([item])
        return shape

    @property
    def clicks(self) -> int:
        return sum(1 for g in self.groups if not g["auto"])

    def _effect(self, ids, item, node_type):
        spid, kind, dur = item["spid"], item["effect"], item["dur"]
        preset, subtype = EFFECTS[kind]
        tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
        body = (f'<p:set><p:cBhvr><p:cTn id="{next(ids)}" dur="1" fill="hold"><p:stCondLst>'
                f'<p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}<p:attrNameLst>'
                f'<p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:to><p:strVal val="visible"/></p:to></p:set>')
        if kind == "zoom":
            for attr in ("ppt_w", "ppt_h"):
                body += (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="{next(ids)}" '
                         f'dur="{dur}" fill="hold"/>{tgt}<p:attrNameLst><p:attrName>{attr}'
                         f'</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
                         f'<p:tav tm="0"><p:val><p:strVal val="#{attr}*0.7"/></p:val></p:tav>'
                         f'<p:tav tm="100000"><p:val><p:strVal val="#{attr}"/></p:val></p:tav>'
                         f'</p:tavLst></p:anim>')
        filt = "wipe(left)" if kind == "wipe" else "fade"
        body += (f'<p:animEffect transition="in" filter="{filt}"><p:cBhvr>'
                 f'<p:cTn id="{next(ids)}" dur="{dur}"/>{tgt}</p:cBhvr></p:animEffect>')
        return (f'<p:par><p:cTn id="{next(ids)}" presetID="{preset}" presetClass="entr" '
                f'presetSubtype="{subtype}" fill="hold" grpId="0" nodeType="{node_type}">'
                f'<p:stCondLst><p:cond delay="{item["delay"]}"/></p:stCondLst>'
                f'<p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')

    def xml(self) -> str:
        counter = iter(range(3, 10_000))
        groups_xml = ""
        for group in self.groups:
            group_id = next(counter)
            steps_xml, offset = "", 0
            for s, step in enumerate(group["steps"]):
                step_id = next(counter)
                effects = ""
                for e, item in enumerate(step):
                    if e > 0:
                        node = "withEffect"
                    elif s == 0 and not group["auto"]:
                        node = "clickEffect"
                    else:
                        node = "afterEffect"
                    effects += self._effect(counter, item, node)
                steps_xml += (f'<p:par><p:cTn id="{step_id}" fill="hold"><p:stCondLst>'
                              f'<p:cond delay="{offset}"/></p:stCondLst><p:childTnLst>'
                              f'{effects}</p:childTnLst></p:cTn></p:par>')
                offset += max(i["delay"] + i["dur"] for i in step)
            start = '<p:cond delay="indefinite"/>'
            if group["auto"]:
                start += '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>'
            groups_xml += (f'<p:par><p:cTn id="{group_id}" fill="hold"><p:stCondLst>{start}'
                           f'</p:stCondLst><p:childTnLst>{steps_xml}</p:childTnLst></p:cTn></p:par>')
        builds = "".join(f'<p:bldP spid="{i}" grpId="0" animBg="1"/>'
                         for i in sorted(self.text_targets))
        bld = f"<p:bldLst>{builds}</p:bldLst>" if builds else ""
        return (f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" '
                f'restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" '
                f'nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
                f'{groups_xml}</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" '
                f'delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
                f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl>'
                f'</p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
                f'{bld}</p:timing>')


def apply_motion(slide, timeline: Timeline):
    """Transición de fundido (spd="med" = 0,75 s) y secuencia de entradas."""
    sld = slide._element
    sld.append(etree.fromstring(f'<p:transition xmlns:p="{P_NS}" spd="med"><p:fade/></p:transition>'))
    if timeline.groups:
        sld.append(etree.fromstring(timeline.xml()))


# ---------------------------------------------------------------------------
# Marco común
# ---------------------------------------------------------------------------

class Deck:
    def __init__(self, assets):
        self.assets = assets
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.notes: list[dict] = []

    def slide(self, dark: bool):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = NOCHE if dark else BLANCO
        return s


def chrome(deck, slide, dark, kicker, title, source, number, title_size=36,
           source_right=W - MARGIN):
    """Encabezado (sección, título, evento) y pie (logo, fuente). Devuelve el grupo del título."""
    fg = BLANCO if dark else AZUL
    accent = DORADO if dark else DORADO_OSCURO
    muted = GRIS_NIEBLA if dark else GRIS
    sh = slide.shapes
    text(sh, W - MARGIN - 5.6, 0.45, 5.6, 0.3, f"{EVENT_TAG} · {number:02d}", 11, muted,
         align=PP_ALIGN.RIGHT, spacing=100, label="evento")
    head = sh.add_group_shape()
    rect(head.shapes, MARGIN, 0.44, 0.05, 0.3, DORADO)
    text(head.shapes, MARGIN + 0.17, 0.45, 6.4, 0.3, kicker.upper(), 13, accent,
         bold=True, spacing=200, label="sección")
    text(head.shapes, MARGIN, 0.80, W - 2 * MARGIN, 0.8, title, title_size, fg,
         bold=True, font=TITLE_FONT, label="título")
    # Pie: divisor fino, logo discreto y fuente.
    rect(sh, MARGIN, 6.80, source_right - MARGIN, 0.012, ACERO if dark else LINEA)
    logo = picture(sh, deck.assets["logo_white" if dark else "logo"], MARGIN, 6.97, h=0.34)
    bar_x = MARGIN + logo.width / 914400 + 0.16
    rect(sh, bar_x, 6.95, 0.025, 0.38, DORADO)
    if source:
        text(sh, bar_x + 0.17, 6.92, source_right - bar_x - 0.17, 0.52, source, 10, muted,
             label="fuente")
    return head


def bubble(slide, x, y, w, h, runs, tail_x, tail_y, size=14.5):
    """Globo de diálogo blanco con colita hacia (tail_x, tail_y) en pulgadas absolutas."""
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = (tail_x - (x + w / 2)) / w
    shp.adjustments[1] = (tail_y - (y + h / 2)) / h
    shp.adjustments[2] = 0.16667
    _plain(shp, BLANCO, LINEA, 1)
    soft_shadow(shp, blur=14, dist=4, alpha=30)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pad = 0.16
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for chunk, keyword in runs:
        run = p.add_run()
        run.text = chunk
        _style_run(run, size, DORADO_OSCURO if keyword else AZUL, True, BODY_FONT)
    measured = [[(c, size, BODY_FONT, True, False) for c, _ in runs]]
    lines = wrapped_lines(measured[0], w - 2 * pad - 0.1)
    if lines > 2:
        FIT_WARNINGS.append(f"globo: «{''.join(c for c, _ in runs)}» ocupa {lines} líneas")
    check_fit("globo", measured, w - 2 * pad - 0.1, h - 0.12)
    return shp


def chillo_peek(deck, slide, x=10.5, width=2.6):
    """Chillo de medio cuerpo asomándose desde el borde inferior derecho."""
    img = Image.open(deck.assets["chillo_half"])
    h = width * img.height / img.width
    return picture(slide.shapes, deck.assets["chillo_half"], x, H - h, w=width)


def peek_bubble(slide, runs):
    """Globo estándar sobre el Chillo asomado: la colita apunta a la cresta."""
    return bubble(slide, CHILLO_RIGHT + 0.18, 2.62, 3.5, 1.0, runs,
                  tail_x=11.75, tail_y=4.1, size=14)


def chillo_full(deck, slide, x, bottom, height):
    return picture(slide.shapes, deck.assets["chillo"], x, bottom - height, h=height)


def card(shapes, x, y, w, h, heading, body, tone="light", heading_size=24,
         body_size=17, number=None, align=PP_ALIGN.LEFT):
    """Tarjeta con filete dorado superior; devuelve el grupo para animarlo junto.

    tone: "light" (niebla sobre blanco), "dark" (acero sobre azul) o "navy" (azul sobre blanco).
    """
    fill, head_c, body_c, num_c = {
        "light": (NIEBLA, AZUL, GRIS, DORADO_OSCURO),
        "dark": (ACERO, BLANCO, GRIS_NIEBLA, DORADO),
        "navy": (AZUL, BLANCO, GRIS_NIEBLA, DORADO),
    }[tone]
    grp = shapes.add_group_shape()
    rect(grp.shapes, x, y, w, h, fill)
    rect(grp.shapes, x, y, w, 0.06, DORADO)
    pad = 0.3
    cx, cw = x + pad, w - 2 * pad
    top = y + 0.32
    head_w = cw
    if number:
        # Número discreto en la esquina; el encabezado le deja sitio y el cuerpo usa todo el ancho.
        text(grp.shapes, x + w - pad - 0.6, top - 0.06, 0.6, 0.5, number, 20, num_c, bold=True,
             font=TITLE_FONT, align=PP_ALIGN.RIGHT, label="número")
        head_w = cw - 0.7
    head_lines = wrapped_lines([(heading, heading_size, TITLE_FONT, True, False)], head_w)
    head_h = head_lines * heading_size * LINE_EM[TITLE_FONT] * 0.92 / 72
    text(grp.shapes, cx, top, head_w, head_h + 0.02, heading, heading_size, head_c,
         bold=True, font=TITLE_FONT, line_spacing=0.92, align=align, label="encabezado")
    body_top = top + head_h + 0.1
    text(grp.shapes, cx, body_top, cw, y + h - body_top - 0.2, body, body_size, body_c,
         line_spacing=1.05, align=align, label="cuerpo")
    return grp


def chip(shapes, x, y, label, fg=NOCHE, bg=DORADO, size=12):
    w = width_in(label, size, BODY_FONT, True) + len(label) * 1.0 / 72 + 0.4
    rect(shapes, x, y, w, 0.34, bg, radius=0.5)
    text(shapes, x, y, w, 0.34, label, size, fg, bold=True, align=PP_ALIGN.CENTER,
         valign=MSO_ANCHOR.MIDDLE, spacing=100, label="chip")
    return w


def record(deck, slide, timeline, title, clicks_detail, chillo=None):
    apply_motion(slide, timeline)
    deck.notes.append({"n": len(deck.notes) + 1, "title": title,
                       "clicks": timeline.clicks, "detail": clicks_detail,
                       "chillo": chillo})


# ---------------------------------------------------------------------------
# Láminas
# ---------------------------------------------------------------------------

def slide_cover(deck):
    s = deck.slide(dark=True)
    sh = s.shapes
    tl = Timeline()
    logo = picture(sh, deck.assets["logo_white"], MARGIN, 0.62, h=1.2)
    tl.add(logo, "fade", dur=500)
    head = sh.add_group_shape()
    rect(head.shapes, MARGIN, 2.62, 0.05, 0.3, DORADO)
    text(head.shapes, MARGIN + 0.17, 2.63, 7.4, 0.3, "MÓDULO DE PLANIFICACIÓN LOGÍSTICA",
         14, DORADO, bold=True, spacing=200)
    text(head.shapes, MARGIN, 3.02, 8.4, 2.15, "Planificar con el costo completo de servir",
         54, BLANCO, bold=True, font=TITLE_FONT, line_spacing=0.95)
    tl.add(head, "wipe", dur=600)
    sub = text(sh, MARGIN, 5.22, 8.6, 0.5,
               [[("Flota propia o contratada. ", {}),
                 ("La IA explica; el planificador decide.", {"color": DORADO_CLARO, "bold": True})]],
               20, GRIS_NIEBLA)
    tl.add(sub, "fade", dur=450)
    foot = sh.add_group_shape()
    rect(foot.shapes, MARGIN, 6.2, 1.1, 0.025, DORADO)
    text(foot.shapes, MARGIN, 6.4, 9.0, 0.35,
         "Equipo IstpetDev · Instituto Superior Tecnológico Traversari", 15, BLANCO)
    text(foot.shapes, MARGIN, 6.82, 9.0, 0.3, EVENT_TAG, 12, GRIS_NIEBLA, spacing=150)
    tl.add(foot, "fade", mode="with", dur=450)
    gallo = chillo_full(deck, s, 9.95, 7.32, 5.05)
    tl.add(gallo, "zoom", dur=550)
    msg = bubble(s, 8.55, 0.72, 4.5, 1.0,
                 [("¡Hola! Soy Chillo y hoy les cuento cómo ", False),
                  ("decidir mejor", True), (" cada ruta.", False)],
                 tail_x=11.2, tail_y=2.35, size=15.5)
    tl.add(msg, "fade", mode="click", dur=400)
    record(deck, s, tl, "Portada",
           "Auto: logo → título → subtítulo y equipo → Chillo (zoom). Clic 1: globo de Chillo.",
           chillo="Chillo entra solo (zoom). Globo: clic 1.")


def slide_problem(deck):
    s = deck.slide(dark=False)
    sh = s.shapes
    tl = Timeline()
    right = CHILLO_RIGHT
    head = chrome(deck, s, False, "Anexo 5.7 · Problema", "¿Flota propia o apoyo contratado?",
                  "Fuente: guion T4, hipótesis de problema A-02; decisión de producto P-02.", 2,
                  source_right=right)
    tl.add(head, "fade", dur=400)
    lead = text(sh, MARGIN, 1.74, right - MARGIN, 0.45,
                "Cada día hay que decidir cómo servir las entregas disponibles.", 21, GRIS)
    tl.add(lead, "fade", mode="with", dur=400)
    gap = 0.3
    cw = (right - MARGIN - gap) / 2
    c1 = card(sh, MARGIN, 2.42, cw, 2.68, "Mirar solo kilómetros",
              "puede ocultar combustible, trabajo de descarga y producto expuesto a merma.",
              heading_size=24, body_size=18)
    tl.add(c1, "wipe", dur=500)
    gallo = chillo_peek(deck, s)
    tl.add(gallo, "zoom", dur=500)
    c2 = card(sh, MARGIN + cw + gap, 2.42, cw, 2.68, "Comparar el servicio completo",
              "ayuda a que el planificador entienda los costos y las entregas pendientes.",
              tone="navy", heading_size=24, body_size=18)
    tl.add(c2, "wipe", mode="click", dur=500)
    hyp = sh.add_group_shape()
    rect(hyp.shapes, MARGIN, 5.38, right - MARGIN, 0.98, CREMA)
    rect(hyp.shapes, MARGIN, 5.38, 0.07, 0.98, DORADO)
    text(hyp.shapes, MARGIN + 0.35, 5.38, right - MARGIN - 0.6, 0.98,
         [[("Nuestra hipótesis: ", {"color": DORADO_OSCURO}),
           ("una decisión explicable entre flota propia y contratada.", {})]],
         20, AZUL, bold=True, valign=MSO_ANCHOR.MIDDLE)
    tl.add(hyp, "fade", dur=450)
    msg = peek_bubble(s, [("Si solo miramos kilómetros, la cuenta sale ", False),
                          ("incompleta", True), (".", False)])
    tl.add(msg, "fade", mode="click", dur=400)
    record(deck, s, tl, "Problema",
           "Auto: título → «Mirar solo kilómetros» (barrido) → Chillo (zoom). "
           "Clic 1: «Comparar el servicio completo» → (auto) hipótesis. Clic 2: globo de Chillo.",
           chillo="Chillo entra solo (zoom). Globo: clic 2.")


def slide_solution(deck):
    s = deck.slide(dark=True)
    sh = s.shapes
    tl = Timeline()
    head = chrome(deck, s, True, "Anexo 5.7 · Solución", "El motor calcula. La persona decide.",
                  "Fuente: matriz A-10, A-11 y P-01/P-02; guion T4.", 3)
    tl.add(head, "fade", dur=400)
    steps = [
        ("01", "Comparar", "Rutas con flota propia y contratada"),
        ("02", "Desglosar", "Costo total estimado de servir"),
        ("03", "Explicar", "La IA comunica los datos calculados"),
        ("04", "Decidir", "El planificador elige la opción"),
    ]
    gap = 0.28
    cw = (W - 2 * MARGIN - 3 * gap) / 4
    d = 0.9
    line = rect(sh, MARGIN + cw / 2, 1.98 + d / 2, W - 2 * MARGIN - cw, 0.02, DORADO)
    tl.add(line, "wipe", dur=500)
    for i, (num, head_txt, body) in enumerate(steps):
        x = MARGIN + i * (cw + gap)
        grp = sh.add_group_shape()
        oval(grp.shapes, x + (cw - d) / 2, 1.98, d, NOCHE, DORADO, 1.5)
        text(grp.shapes, x + (cw - d) / 2, 1.98, d, d, num, 24, DORADO, bold=True,
             font=TITLE_FONT, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        rect(grp.shapes, x, 3.12, cw, 1.62, ACERO)
        rect(grp.shapes, x, 3.12, cw, 0.05, DORADO)
        text(grp.shapes, x + 0.2, 3.3, cw - 0.4, 0.55, head_txt, 26, BLANCO, bold=True,
             font=TITLE_FONT, align=PP_ALIGN.CENTER)
        text(grp.shapes, x + 0.2, 3.9, cw - 0.4, 0.75, body, 16, GRIS_NIEBLA,
             align=PP_ALIGN.CENTER, line_spacing=1.0)
        tl.add(grp, "fade", mode="click", dur=450)
    tag = text(sh, MARGIN, 5.02, W - 2 * MARGIN, 0.35,
               "PROPUESTA · la comparación completa A/B/C sigue en integración",
               13, DORADO_CLARO, bold=True, align=PP_ALIGN.CENTER, spacing=120)
    tl.add(tag, "fade", dur=400)
    claim = sh.add_group_shape()
    claim_runs = [("La IA explica; ", {}), ("el planificador decide.", {"color": DORADO})]
    tw = sum(width_in(c, 32, TITLE_FONT, True) for c, _ in claim_runs)
    text(claim.shapes, (W - tw - 0.4) / 2, 5.55, tw + 0.4, 0.8, [claim_runs], 32, BLANCO,
         bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    rect(claim.shapes, (W - tw) / 2 - 1.55, 5.96, 1.2, 0.02, DORADO)
    rect(claim.shapes, (W + tw) / 2 + 0.35, 5.96, 1.2, 0.02, DORADO)
    tl.add(claim, "fade", dur=450)
    record(deck, s, tl, "Solución",
           "Auto: título → línea dorada. Clic 1: Comparar. Clic 2: Desglosar. Clic 3: Explicar. "
           "Clic 4: Decidir → (auto) rótulo PROPUESTA → «La IA explica; el planificador decide».")


def slide_capture(deck):
    s = deck.slide(dark=False)
    sh = s.shapes
    tl = Timeline()
    head = chrome(deck, s, False, "Anexo 5.7 · Prototipo / demo",
                  "La pantalla de rutas, en una captura real",
                  "Fuente: captura-rutas.png, Entregable 2. La captura muestra otro escenario "
                  "sintético; no es Quito ni A/B/C.", 4)
    tl.add(head, "fade", dur=400)
    frame = sh.add_group_shape()
    img = Image.open(deck.assets["capture"])
    pw = 7.7
    ph = pw * img.height / img.width
    top = 1.85
    border = rect(frame.shapes, MARGIN, top - 0.08, pw + 0.16, ph + 0.16, BLANCO, LINEA)
    soft_shadow(border, blur=12, dist=4, alpha=18)
    picture(frame.shapes, deck.assets["capture"], MARGIN + 0.08, top, w=pw)
    rect(frame.shapes, MARGIN, top + ph + 0.08, pw + 0.16, 0.05, DORADO)
    tl.add(frame, "wipe", dur=600)
    x2 = MARGIN + pw + 0.5
    info = card(sh, x2, top - 0.08, W - MARGIN - x2, ph + 0.21, "Interfaz existente",
                "La captura documenta una vista de rutas. No demuestra la comparación de flota "
                "ni una corrida del escenario de Quito.", heading_size=26, body_size=17)
    tl.add(info, "fade", dur=450)
    label = sh.add_group_shape()
    y = top + ph + 0.55
    rect(label.shapes, MARGIN, y, W - 2 * MARGIN, 0.82, AZUL)
    rect(label.shapes, MARGIN, y, 0.07, 0.82, DORADO)
    cw = chip(label.shapes, MARGIN + 0.35, y + 0.24, "DATOS SINTÉTICOS")
    text(label.shapes, MARGIN + 0.35 + cw + 0.3, y, W - 2 * MARGIN - cw - 0.9, 0.82,
         "Captura real · los datos visibles pertenecen a otro escenario sintético.",
         18, BLANCO, bold=True, valign=MSO_ANCHOR.MIDDLE)
    tl.add(label, "fade", dur=450)
    record(deck, s, tl, "Prototipo / demo · captura",
           "Todo automático: título → captura (barrido) → tarjeta → rótulo «datos sintéticos».")


# Cifras de la columna A: selector de #240 (PR #254, merge 6adcf53), fixture demo-base,
# BOD-UIO -> HOS-001 -> regreso. Fuente: 003-agente-logistico/evidencia/240-desglose-demo.json
# (match.serviceCostBreakdown). B y C dependen de #241 y quedan en integración (matriz P-02).
ABC_ENGINE_COMMIT = "6adcf53"
ABC_COLUMN_A = ["USD 2,26", "USD 2,15", "USD 7,22 + 4,81", "USD 3,50", "USD 19,94"]
ABC_PENDING = "en integración"


def slide_abc(deck):
    s = deck.slide(dark=False)
    sh = s.shapes
    tl = Timeline()
    right = CHILLO_RIGHT
    head = chrome(deck, s, False, "Anexo 5.7 · Prototipo / demo",
                  "A / B / C: la columna A ya la calcula el motor",
                  "Fuente: #240, PR #254 (6adcf53), 240-desglose-demo.json · demo-base: Bodega Quito → "
                  "Hospital Norte → regreso, 210 kg, 20 bultos, 1 ayudante · combustible sintético "
                  "USD 3,785/gal · matriz P-01; P-02 en curso.", 5, source_right=right)
    tl.add(head, "fade", dur=400)
    band = sh.add_group_shape()
    cw = chip(band.shapes, MARGIN, 1.68, "DATOS SINTÉTICOS")
    text(band.shapes, MARGIN + cw + 0.2, 1.68, right - MARGIN - cw - 0.2, 0.34,
         f"A: costo estimado del motor · escenario demo-base · commit {ABC_ENGINE_COMMIT}",
         13, GRIS, bold=True, valign=MSO_ANCHOR.MIDDLE)
    tl.add(band, "fade", mode="with", dur=400)

    x0, y0 = MARGIN, 2.2
    widths = [2.7, 2.45, 1.8, 1.8]
    headings = ["Componente", "A · Flota propia", "B · Propia, críticos primero", "C · Contratada"]
    header = sh.add_group_shape()
    xx = x0
    for j, (w, title) in enumerate(zip(widths, headings)):
        fill = NIEBLA if j == 0 else (AZUL if j == 1 else GRIS)
        color = AZUL if j == 0 else BLANCO
        rect(header.shapes, xx, y0, w, 0.66, fill)
        text(header.shapes, xx + 0.1, y0, w - 0.2, 0.66, title, 15, color, bold=True,
             font=TITLE_FONT, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE,
             line_spacing=0.9)
        xx += w
    rect(header.shapes, x0 + widths[0], y0 + 0.62, widths[1], 0.04, DORADO)
    tl.add(header, "fade", dur=400)
    gallo = chillo_peek(deck, s)
    tl.add(gallo, "zoom", dur=500)

    labels = ["Combustible", "Desgaste", "Conductor y ayudantes", "Merma esperada", "Costo total"]
    rh = 0.47
    rows = []
    yy = y0 + 0.66
    for i, label in enumerate(labels):
        total = i == len(labels) - 1
        h = rh + (0.12 if total else 0)
        grp = sh.add_group_shape()
        xx = x0
        for j, w in enumerate(widths):
            fill = CREMA if total else (BLANCO if i % 2 == 0 else NIEBLA)
            rect(grp.shapes, xx, yy, w, h, fill)
            if j == 0:
                text(grp.shapes, xx + 0.2, yy, w - 0.25, h, label, 14, AZUL, bold=True,
                     valign=MSO_ANCHOR.MIDDLE)
            elif j == 1:
                text(grp.shapes, xx + 0.1, yy, w - 0.2, h, ABC_COLUMN_A[i], 25 if total else 19,
                     AZUL, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER,
                     valign=MSO_ANCHOR.MIDDLE)
            else:
                text(grp.shapes, xx + 0.1, yy, w - 0.2, h, ABC_PENDING, 13, GRIS, italic=True,
                     align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
            xx += w
        rect(grp.shapes, x0, yy + h - 0.012, sum(widths), 0.012, LINEA)
        if total:
            rect(grp.shapes, x0, yy, sum(widths), 0.03, DORADO)
        rows.append(grp)
        yy += h
    for i, grp in enumerate(rows[:-1]):
        tl.add(grp, "wipe", mode="click" if i == 0 else "after", dur=300)
    tl.add(rows[-1], "fade", mode="click", dur=450)

    quote = text(sh, MARGIN, yy + 0.14, right - MARGIN, 0.6,
                 "«El costo de una ruta no es solo kilómetros: sumamos combustible según carga, "
                 "personas que descargan y merma esperada».",
                 14, AZUL, italic=True)
    tl.add(quote, "fade", dur=400)
    notes = sh.add_group_shape()
    ny = 6.08
    left_w = 5.55
    rect(notes.shapes, MARGIN, ny, left_w, 0.6, NIEBLA)
    rect(notes.shapes, MARGIN, ny, 0.06, 0.6, DORADO)
    text(notes.shapes, MARGIN + 0.22, ny, left_w - 0.35, 0.6,
         "La comparación con operador contratado y la recomendación calculada se integran en #241.",
         12, AZUL, bold=True, valign=MSO_ANCHOR.MIDDLE)
    px = MARGIN + left_w + 0.15
    rect(notes.shapes, px, ny, right - px, 0.6, BLANCO, GRIS, line_width=1)
    text(notes.shapes, px + 0.1, ny, right - px - 0.2, 0.6,
         "[CAPTURA PENDIENTE: pantalla A/B/C]", 12, GRIS, bold=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    tl.add(notes, "fade", mode="with", dur=400)
    msg = peek_bubble(s, [("Una ruta no es solo kilómetros: también es ", False),
                          ("quien descarga", True), (".", False)])
    tl.add(msg, "fade", mode="click", dur=400)
    record(deck, s, tl, "Panel A / B / C",
           "Auto: título y rótulo → encabezado de tabla → Chillo (zoom). Clic 1: filas Combustible, "
           "Desgaste, Conductor y ayudantes, Merma (en cascada). Clic 2: Costo total → (auto) frase "
           "y notas #241 / captura pendiente. Clic 3: globo de Chillo.",
           chillo="Chillo entra solo (zoom). Globo: clic 3.")


def slide_impact(deck):
    s = deck.slide(dark=True)
    sh = s.shapes
    tl = Timeline()
    head = chrome(deck, s, True, "Anexo 5.7 · Impacto", "Indicadores que cualquier rol puede entender",
                  "Propuesta de indicadores: matriz P-07 en curso; diseño de medición según la "
                  "propuesta de planificación.", 6)
    text(head.shapes, MARGIN, 1.66, 10.0, 0.4,
         "Preguntas sencillas para conversar con el equipo operativo", 19, GRIS_NIEBLA)
    tl.add(head, "fade", dur=400)
    items = [
        ("Costo por kilómetro", "¿Cuánto cuesta recorrer esta ruta?"),
        ("Costo por entrega", "¿Cuánto cuesta atender cada pedido?"),
        ("Entregas a tiempo", "¿Llegamos dentro de la ventana acordada?"),
        ("Pendientes", "¿Qué pedidos quedaron sin atender?"),
        ("Mermas", "¿Qué producto se dañó o rechazó?"),
        ("Propia frente a contratada", "¿Qué opción conviene con los datos disponibles?"),
    ]
    gap = 0.3
    cw = (W - 2 * MARGIN - 2 * gap) / 3
    for i, (head_txt, body) in enumerate(items):
        x = MARGIN + (i % 3) * (cw + gap)
        y = 2.28 + (i // 3) * 2.24
        grp = card(sh, x, y, cw, 2.08, head_txt, body, tone="dark", heading_size=21,
                   body_size=16, number=f"{i + 1:02d}")
        tl.add(grp, "fade", dur=350)
    record(deck, s, tl, "Impacto",
           "Todo automático: título → seis indicadores en cascada.")


def slide_viability(deck):
    s = deck.slide(dark=False)
    sh = s.shapes
    tl = Timeline()
    head = chrome(deck, s, False, "Anexo 5.7 · Viabilidad", "Lo que podemos respaldar hoy",
                  "Fuente: matriz P-01/P-04/P-11 y A-22. Ensayo P-01: sintético, fixture demo-base; "
                  "no es la demo A/B/C.", 7)
    tl.add(head, "fade", dur=400)
    gap = 0.35
    cw = (W - 2 * MARGIN - gap) / 2
    c1 = card(sh, MARGIN, 1.95, cw, 2.3, "Costo estimado de servir",
              "Ensayo técnico sintético: combustible según carga, desgaste, conductor, ayudantes y "
              "merma esperada.", number="01", heading_size=26, body_size=18)
    tl.add(c1, "fade", dur=450)
    c2 = card(sh, MARGIN + cw + gap, 1.95, cw, 2.3, "Importación y escenario",
              "Hay prueba de importación del escenario sintético de Quito. La comparación completa "
              "no está ensayada en pantalla.", number="02", heading_size=26, body_size=18)
    tl.add(c2, "fade", dur=450)
    lim = sh.add_group_shape()
    rect(lim.shapes, MARGIN, 4.75, W - 2 * MARGIN, 1.35, AZUL)
    rect(lim.shapes, MARGIN, 4.75, 0.07, 1.35, DORADO)
    text(lim.shapes, MARGIN + 0.45, 4.75, 1.4, 1.35, "LÍMITE", 22, DORADO, bold=True,
         font=TITLE_FONT, valign=MSO_ANCHOR.MIDDLE, spacing=150)
    text(lim.shapes, MARGIN + 1.95, 4.75, W - 2 * MARGIN - 2.3, 1.35,
         "Costo estimado y pruebas del motor no equivalen a resultados medidos de una empresa.",
         22, BLANCO, bold=True, valign=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    tl.add(lim, "wipe", dur=500)
    record(deck, s, tl, "Viabilidad",
           "Todo automático: título → tarjeta 01 → tarjeta 02 → banda «Límite» (barrido).")


def slide_business(deck):
    s = deck.slide(dark=True)
    sh = s.shapes
    tl = Timeline()
    head = chrome(deck, s, True, "Anexo 5.7 · Modelo de negocio", "Una propuesta comercial por validar",
                  "Fuente: matriz A-26/A-30/A-31. Sin precio, ventas ni costos mensuales confirmados.", 8)
    tl.add(head, "fade", dur=400)
    gap = 0.35
    cw = (W - 2 * MARGIN - gap) / 2
    c1 = card(sh, MARGIN, 1.95, cw, 2.3, "Cliente propuesto",
              "Empresas de servicios de limpieza con operaciones de abastecimiento y rutas por "
              "planificar.", tone="dark", number="01", heading_size=28, body_size=19)
    tl.add(c1, "fade", dur=450)
    c2 = card(sh, MARGIN + cw + gap, 1.95, cw, 2.3, "Cómo se propone cobrar",
              "Implementación e integración inicial, más una suscripción. Precio y disposición de "
              "pago por validar.", tone="dark", number="02", heading_size=28, body_size=19)
    tl.add(c2, "fade", dur=450)
    closing = sh.add_group_shape()
    rect(closing.shapes, (W - 1.1) / 2, 4.82, 1.1, 0.025, DORADO)
    text(closing.shapes, MARGIN, 5.05, W - 2 * MARGIN, 0.9,
         "El piloto serviría para validar el uso, los costos reales y el valor para el cliente.",
         22, BLANCO, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    tl.add(closing, "fade", dur=450)
    record(deck, s, tl, "Modelo de negocio",
           "Todo automático: título → tarjeta 01 → tarjeta 02 → frase del piloto.")


def slide_risks(deck):
    s = deck.slide(dark=False)
    sh = s.shapes
    tl = Timeline()
    right = CHILLO_RIGHT
    head = chrome(deck, s, False, "Anexo 5.7 · Riesgos y mitigación",
                  "RIESGOS Y ESTRATEGIAS DE MITIGACIÓN",
                  "Fuente: matriz P-16 y diseño de riesgos de la propuesta del evento. Medidas "
                  "propuestas, no todas implementadas.", 9, source_right=right)
    tl.add(head, "fade", dur=400)
    gallo = chillo_peek(deck, s)
    tl.add(gallo, "zoom", dur=500)
    rows = [
        ("Datos incompletos", "Validar al importar; mostrar ‘sin dato’ y supuestos."),
        ("Ruta poco realista", "Revisar ventanas, capacidad y restricciones antes de publicar."),
        ("Costo estimado inexacto", "Mostrar el desglose y calibrar con viajes conciliados."),
        ("Error en la explicación de IA", "Basarla en cifras calculadas; el planificador decide."),
        ("Resistencia al cambio", "Capacitar por rol y revisar el plan con el equipo."),
    ]
    rw = 3.25
    mx = MARGIN + rw + 0.45
    y, rh = 1.72, 0.74
    for risk, mitigation in rows:
        grp = sh.add_group_shape()
        rect(grp.shapes, MARGIN, y, rw, rh, AZUL)
        text(grp.shapes, MARGIN + 0.22, y, rw - 0.3, rh, risk, 17, BLANCO, bold=True,
             font=TITLE_FONT, valign=MSO_ANCHOR.MIDDLE)
        arrow = grp.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(MARGIN + rw + 0.13),
                                     Inches(y + rh / 2 - 0.09), Inches(0.2), Inches(0.18))
        arrow.rotation = 90
        _plain(arrow, DORADO, None, 0)
        rect(grp.shapes, mx, y, right - mx, rh, NIEBLA)
        text(grp.shapes, mx + 0.25, y, right - mx - 0.4, rh, mitigation, 15.5, AZUL,
             valign=MSO_ANCHOR.MIDDLE)
        tl.add(grp, "wipe", dur=350)
        y += rh + 0.13
    note = text(sh, MARGIN, 6.14, right - MARGIN, 0.4,
                "La charla de logística destacó especialmente este análisis.", 14, DORADO_OSCURO,
                italic=True, bold=True)
    tl.add(note, "fade", dur=400)
    msg = peek_bubble(s, [("Todo plan tiene riesgos; lo clave es saber ", False),
                          ("mitigarlos", True), (".", False)])
    tl.add(msg, "fade", mode="click", dur=400)
    record(deck, s, tl, "Riesgos y mitigación",
           "Auto: título → Chillo (zoom) → cinco riesgos en cascada → nota de la charla. "
           "Clic 1: globo de Chillo.",
           chillo="Chillo entra solo (zoom). Globo: clic 1.")


def slide_roadmap(deck):
    s = deck.slide(dark=True)
    sh = s.shapes
    tl = Timeline()
    right = CHILLO_RIGHT
    head = chrome(deck, s, True, "Anexo 5.7 · Próximo paso", "Hoja de ruta: validar antes de prometer",
                  "Fuente: matriz P-02/P-05/P-10/P-12/P-13/P-14/P-15 y A-32. Propuestas sujetas a "
                  "validación.", 10, source_right=right)
    tl.add(head, "fade", dur=400)
    gallo = chillo_full(deck, s, 10.42, 6.72, 4.3)
    tl.add(gallo, "zoom", dur=550)
    steps = [
        ("Integrar", "Comparación propia / contratada y recomendación justificada"),
        ("Validar Quito", "Tráfico y restricciones con fuentes vigentes"),
        ("Completar el ciclo", "Guía de retorno para entregas pendientes"),
        ("Calibrar", "Pronóstico y costos con datos conciliados"),
    ]
    gap = 0.28
    cw = (right - MARGIN - gap) / 2
    for i, (head_txt, body) in enumerate(steps):
        x = MARGIN + (i % 2) * (cw + gap)
        y = 1.78 + (i // 2) * 1.76
        grp = card(sh, x, y, cw, 1.64, head_txt, body, tone="dark", heading_size=22,
                   body_size=15.5, number=f"{i + 1:02d}")
        tl.add(grp, "fade", dur=350)
    pilot = sh.add_group_shape()
    rect(pilot.shapes, MARGIN, 5.35, right - MARGIN, 1.15, DORADO)
    text(pilot.shapes, MARGIN + 0.35, 5.35, right - MARGIN - 0.7, 1.15,
         [[("Próximo paso: ", {"font": TITLE_FONT, "size": 23}),
           ("acordar un piloto, una línea base y los indicadores con una empresa.", {})]],
         18, NOCHE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    tl.add(pilot, "wipe", mode="click", dur=550)
    msg = bubble(s, CHILLO_RIGHT + 0.18, 1.3, 3.5, 1.0,
                 [("¿Lo validamos juntos en un ", False), ("piloto", True),
                  ("? ¡Kikirikí!", False)],
                 tail_x=11.55, tail_y=2.48, size=14)
    tl.add(msg, "fade", mode="click", dur=400)
    record(deck, s, tl, "Hoja de ruta y cierre",
           "Auto: título → Chillo de cuerpo entero (zoom) → cuatro pasos en cascada. "
           "Clic 1: banda dorada «Próximo paso» (barrido). Clic 2: globo de Chillo.",
           chillo="Chillo entra solo (zoom). Globo: clic 2.")


# ---------------------------------------------------------------------------
# Construcción y exportación
# ---------------------------------------------------------------------------

def export(pptx: Path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("No se encontró soffice/libreoffice para generar el PDF")
    if PDF_PATH.exists():
        PDF_PATH.unlink()
    profile = Path("/tmp/lo-profile-pitch-v3")
    profile.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless",
         "--convert-to", "pdf", "--outdir", str(HERE), str(pptx)],
        check=True, capture_output=True, text=True,
    )
    if not PDF_PATH.exists():
        raise RuntimeError(f"LibreOffice no produjo el PDF: {result.stdout} {result.stderr}")
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise RuntimeError("No se encontró pdftoppm para renderizar los PNG")
    RENDER_DIR.mkdir(exist_ok=True)
    for old in RENDER_DIR.glob("diapositiva-*.png"):
        old.unlink()
    # 110 ppp equivale a 1467×825 px para 13,333×7,5 in, igual que render-v2.
    subprocess.run([pdftoppm, "-png", "-r", "110", str(PDF_PATH),
                    str(RENDER_DIR / "diapositiva")], check=True)
    # pdftoppm rellena con ceros a partir de 10 páginas; se normaliza a diapositiva-N.png.
    for png in RENDER_DIR.glob("diapositiva-*.png"):
        png.rename(RENDER_DIR / f"diapositiva-{int(png.stem.split('-')[1])}.png")


def build_deck():
    for required in (ROUTE_CAPTURE, ASSETS / "chillo-el-gallo.png", ASSETS / "ISTPET-ORIGINAL.png",
                     FONTS / "Oswald.ttf", FONTS / "OpenSans.ttf", FONTS / "OpenSans-Italic.ttf"):
        if not required.exists():
            raise FileNotFoundError(f"Falta un recurso requerido: {required}")
    deck = Deck(prepare_assets())
    deck.prs.core_properties.title = "IstpetDev · Pitch Expo Clean · v3"
    deck.prs.core_properties.subject = "Planificación logística, basado en guion T4"
    for builder in (slide_cover, slide_problem, slide_solution, slide_capture, slide_abc,
                    slide_impact, slide_viability, slide_business, slide_risks, slide_roadmap):
        builder(deck)
    deck.prs.save(PPTX_PATH)
    # Reapertura de control: falla aquí si el XML inyectado deja el paquete ilegible.
    reopened = Presentation(str(PPTX_PATH))
    assert len(reopened.slides) == len(deck.notes)
    export(PPTX_PATH)
    print(f"PPTX: {PPTX_PATH}")
    print(f"PDF:  {PDF_PATH}")
    print(f"PNG:  {RENDER_DIR}")
    for note in deck.notes:
        print(f"{note['n']:>2}. {note['title']}: {note['clicks']} clic(s) · {note['detail']}")
    for warning in FIT_WARNINGS:
        print(f"AVISO de ajuste · {warning}")


if __name__ == "__main__":
    build_deck()
