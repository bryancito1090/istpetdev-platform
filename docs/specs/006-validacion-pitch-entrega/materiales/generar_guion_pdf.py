"""Genera Entregable/8-Guion-del-pitch.pdf desde la tabla de la «Propuesta v7» de la nota
docs/boveda/07 Demo/Guion del pitch.md (la nota es la fuente; este PDF se regenera).

Uso:
    uv run --with python-docx python docs/specs/006-validacion-pitch-entrega/materiales/generar_guion_pdf.py
"""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[4]
NOTE = ROOT / "docs/boveda/07 Demo/Guion del pitch.md"
OUT = ROOT / "docs/boveda/Entregable/8-Guion-del-pitch.pdf"
AZUL, DORADO = RGBColor(0x24, 0x2D, 0x57), RGBColor(0x7D, 0x65, 0x20)


def rows():
    text = NOTE.read_text(encoding="utf-8")
    section = text.split("## Propuesta v7", 1)[1].split("\n---\n", 1)[0]
    for line in section.splitlines():
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) >= 5 and cells[0].isdigit():
            yield int(cells[0]), cells[2], int(cells[3]), cells[4].strip("«»")


def stamp(seconds: int) -> str:
    return f"{seconds // 60:02d}:{seconds % 60:02d}"


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Open Sans", Pt(11)
    items = list(rows())
    total = sum(s for _, _, s, _ in items)
    title = doc.add_paragraph()
    run = title.add_run("Guion del pitch v7 — IstpetDev")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(20), AZUL
    sub = doc.add_paragraph(f"{total // 60} minutos {total % 60} segundos · {len(items)} diapositivas "
                            "· Reto 1, Expo Clean 2026 · Optimización y trazabilidad logística")
    sub.runs[0].font.color.rgb = DORADO
    t = 0
    for n, voice, secs, spoken in items:
        head = doc.add_paragraph()
        r = head.add_run(f"{stamp(t)}–{stamp(t + secs)} · {voice} · Diapositiva {n}")
        r.bold, r.font.color.rgb = True, AZUL
        head.paragraph_format.space_before = Pt(10)
        head.paragraph_format.space_after = Pt(2)
        doc.add_paragraph(re.sub(r"\s+", " ", spoken))
        t += secs
    with tempfile.TemporaryDirectory() as tmp:
        docx = Path(tmp) / "8-Guion-del-pitch.docx"
        doc.save(docx)
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        subprocess.run([soffice, f"-env:UserInstallation=file://{tmp}/lo", "--headless",
                        "--convert-to", "pdf", "--outdir", tmp, str(docx)], check=True,
                       capture_output=True)
        shutil.copy(Path(tmp) / "8-Guion-del-pitch.pdf", OUT)
    print(f"{OUT} · {len(items)} diapositivas · {total} s")


if __name__ == "__main__":
    main()
