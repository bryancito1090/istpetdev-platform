"""Genera el informe de avance del Checkpoint del 16-oct (evolución de la idea) en Word.

Uso: python3 generar.py  (desde esta carpeta). Luego convertir a PDF con LibreOffice:
    soffice --headless --convert-to pdf IstpetDev-Informe-Checkpoint-16oct.docx
Mismo estilo que el entregable 2 (../entregable-2/generar.py). Fuentes: bóveda (NotasEvento,
02 Producto, 06 Equipo) y PR fusionados en develop el 16-oct-2026.
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
LOGO = AQUI.parent / "entregable-2" / "img" / "logo.png"
CYAN, LIME, ORANGE, DARK, GREY = "00A7C4", "9EC200", "EA601F", "2D3748", "64748B"


def rgb(h):
    return RGBColor.from_string(h)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def borders(table, color="CBD5E1", size="4"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single" if size != "0" else "nil")
        e.set(qn("w:sz"), size)
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def left_bar(paragraph, color, size="36"):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    bdr.append(left)
    pPr.append(bdr)


def bottom_rule(paragraph, color=CYAN, size="12"):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), size)
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), color)
    bdr.append(b)
    pPr.append(bdr)


def fix_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips)))
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        for c, w in zip(row.cells, widths):
            c.width = w


def vcenter(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), "center")
    tcPr.append(v)


def para(doc_or_cell, text="", size=10.5, bold=False, color=DARK, align=None, after=4, italic=False):
    p = doc_or_cell.add_paragraph()
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = rgb(color)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    return p


def rich(p, parts, size=10.5):
    """parts: lista de (texto, negrita)."""
    for text, b in parts:
        r = p.add_run(text)
        r.bold = b
        r.font.size = Pt(size)
        r.font.color.rgb = rgb(DARK)
    return p


def h1(doc, num, text):
    p = doc.add_heading(level=1)
    r = p.add_run(f"{num}  ")
    r.font.color.rgb = rgb(CYAN)
    r2 = p.add_run(text)
    r2.font.color.rgb = rgb(DARK)
    for run in (r, r2):
        run.font.size = Pt(17)
        run.bold = True
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    bottom_rule(p)
    return p


def bullet(doc, parts, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    rich(p, parts, size)
    p.paragraph_format.space_after = Pt(3)
    return p


def chip(cell, text, color, size=8.5):
    cell.text = ""
    shade(cell, color)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = rgb("FFFFFF")
    vcenter(cell)


def cell_text(cell, text, size=9.5, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = rgb(DARK)
    return p


def header_row(table, labels, widths):
    hdr = table.rows[0].cells
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))
    for c, txt, w in zip(hdr, labels, widths):
        c.width = w
        chip(c, txt, DARK)


def footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("IstpetDev · Hackathon Expo Clean Ecuador 2026 · Reto 1 · Informe de avance — Checkpoint del 16-oct · pág. ")
    r.font.size = Pt(8)
    r.font.color.rgb = rgb(GREY)
    run = p.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), tag)
            run._r.append(f)
        else:
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = txt
            run._r.append(t)
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(GREY)


def callout(doc, parts, border=CYAN, fill="E0F7FB", size=10.5):
    box = doc.add_table(rows=1, cols=1)
    borders(box, border, "8")
    cell = box.cell(0, 0)
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    rich(p, parts, size)
    fix_widths(box, (Cm(16.6),))
    para(doc, "", after=4)
    return cell


doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(10.5)
st.font.color.rgb = rgb(DARK)
for s in doc.sections:
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.top_margin = s.bottom_margin = Cm(2.0)
    s.left_margin = s.right_margin = Cm(2.2)

# ---------------- PORTADA ----------------
cover = doc.sections[0]
cover.different_first_page_header_footer = True
para(doc, "", after=40)
doc.add_picture(str(LOGO), width=Cm(11))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
para(doc, "", after=50)
para(doc, "HACKATHON EXPO CLEAN ECUADOR 2026", 11, True, CYAN, after=2)
para(doc, "Reto 1 — Optimización y trazabilidad logística del abastecimiento de insumos", 11, False, GREY, after=26)
para(doc, "Evolución de la idea", 30, True, DARK, after=0)
para(doc, "durante el día 1", 30, True, CYAN, after=14)
p = para(doc, "Informe de avance · Checkpoint de avances del 16 de octubre", 12, False, DARK, after=60)
left_bar(p, ORANGE, "24")

tb = doc.add_table(rows=1, cols=2)
tb.alignment = WD_TABLE_ALIGNMENT.LEFT
borders(tb, "FFFFFF", "0")
c1, c2 = tb.rows[0].cells
c1.width, c2.width = Cm(8.3), Cm(8.3)
para(c1, "EQUIPO", 9, True, GREY, after=2)
para(c1, "IstpetDev", 13, True, DARK, after=8)
for name in ("Bryan Baño", "Jorge Doicela", "Joseph Baño", "Anthony Pilatasig", "Alexander Martínez"):
    para(c1, name, 10.5, False, DARK, after=1)
para(c2, "FECHA", 9, True, GREY, after=2)
para(c2, "16 de octubre de 2026, 18:00", 11, False, DARK, after=8)
para(c2, "LUGAR", 9, True, GREY, after=2)
para(c2, "Quito, Ecuador", 11, False, DARK, after=8)
para(c2, "REVISAN", 9, True, GREY, after=2)
para(c2, "Mónica Sevilla y William Chamorro (Checkpoint 18:30–19:00)", 11, False, DARK, after=8)
for c in (c1, c2):
    c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)

# ---------------- CUERPO ----------------
body = doc.add_section(WD_SECTION.NEW_PAGE)
footer(body)

h1(doc, "1", "Resumen")
callout(doc, [
    ("Llegamos ", True), ("con una plataforma para que una empresa de limpieza reabastezca sus puntos antes de que falte un insumo, "
                          "planifique rutas y demuestre cada entrega con QR, incluso sin señal. ", False),
    ("Hoy escuchamos ", True), ("al sector en dos charlas, preguntamos por sus dolores y preparamos la mentoría técnica. ", False),
    ("Aprendimos ", True), ("que el problema más caro está en la planificación: a quién atender primero, cuánto cuesta de verdad cada ruta "
                            "y cuándo conviene contratar logística porque la flota propia no alcanza. ", False),
    ("Por eso reorientamos ", True), ("el producto hacia un módulo de planificación que compara opciones, suma todo lo que cuesta servir "
                                      "y recomienda con justificación, ", False),
    ("e integramos hoy al menos 14 cambios funcionales ", True), ("en el prototipo.", False),
])

h1(doc, "2", "La idea con la que llegamos")
para(doc, "Nuestra propuesta inicial se organizaba en tres momentos. Este informe describe cómo cambió durante el día 1.", after=6)
prev = doc.add_table(rows=1, cols=2)
borders(prev, "E2E8F0", "4")
w = (Cm(3.2), Cm(13.4))
header_row(prev, ("Momento", "Qué proponíamos"), w)
for etapa, color, texto in [
    ("Anticipar", CYAN, "Importar catálogo, puntos, stock y consumo; calcular los días de cobertura de cada punto frente al tiempo de reposición y clasificar el riesgo."),
    ("Decidir", LIME, "Proponer bodega de origen, vehículo y ruta por calles reales respetando ventanas y jornada; un agente (n8n + DeepSeek) explica y una persona aprueba."),
    ("Demostrar", ORANGE, "Custodia de cada caja con QR desde la bodega hasta la recepción, con una aplicación de campo que funciona sin conexión."),
]:
    row = prev.add_row().cells
    chip(row[0], etapa.upper(), color)
    cell_text(row[1], texto)
fix_widths(prev, w)

h1(doc, "3", "Lo que escuchamos y preguntamos hoy")
para(doc, "Registramos cada intervención en notas colaborativas por participante dentro de nuestra bóveda de documentación. "
          "Las frases son paráfrasis salvo que se indique lo contrario; las cifras oídas no se usan como datos verificados.", 9.5, False, GREY, after=6)
src = doc.add_table(rows=1, cols=3)
borders(src, "E2E8F0", "4")
w = (Cm(3.6), Cm(7.2), Cm(5.8))
header_row(src, ("Fuente", "Qué aprendimos", "Qué decidimos"), w)
for fuente, aprendimos, decidimos in [
    ("Charla del sector limpieza (mañana)",
     "Las empresas abastecen con fichas semanales o quincenales, pero hay clientes que no programan compras. Un catálogo puede llegar a unos 400 ítems; las cajas se abren y se subdividen. "
     "Se consolidan muchas entregas por salida y se reparte por corredores y zonas. Un quirófano no se abastece igual que una habitación.",
     "Priorizar por criticidad y hora límite; conservar lo que no cabe en el camión como pendiente; preparar los datos por zona y presentación."),
    ("Preguntas del equipo a la oradora",
     "Los dolores del servicio se concentran en el personal: capacitación, dosificación, aplicación correcta y control de calidad. "
     "Para personal especializado hay capacitación y protocolos de emergencia.",
     "Agregar protocolos de capacitación por rol e indicadores que entienda toda la empresa."),
    ("Charla de logística integral (tarde)",
     "La logística fragmentada erosiona margen y servicio. Hay que mirar almacén, transporte y punto de servicio a la vez; "
     "medir costo por km, faltantes, servicio, tiempo, mermas y exactitud; y calcular el costo de corregir una entrega. "
     "Una empresa contrata logística cuando su flota propia no alcanza o no conviene.",
     "Cambiar el foco al módulo de planificación: flota propia o contratada, costo total de servir y riesgos con mitigación."),
    ("Diapositivas de esa charla",
     "Riesgos y mitigación: datos inconsistentes, resistencia al cambio, rutas poco realistas, pronósticos imprecisos («medir el error»), "
     "falta de integración y costos de implementación («comenzar con un MVP con beneficios medibles»). Ejemplo: ventana de recepción 07:00–09:30.",
     "Incorporar una sección de riesgos y mitigación al pitch y medir el error del pronóstico."),
    ("Mentoría técnica asignada (Mauro Zurita, 16:15–17:00)",
     "Llevamos cinco preguntas: qué entrega va primero cuando no alcanza el camión, si nuestra regla de reposición se parece a la de la industria, "
     "si comparar contra un pedido fijo cada 7 días es justo, qué prueba de recepción acepta el cliente y qué indicador miraría un gerente para pagar.",
     "Las respuestas se contrastan en la validación integral del sábado con todos los mentores."),
    ("Notas internas del equipo",
     "El planificador necesita ajustar peso y volumen de productos, calibrar costos de vehículos, agrupar por zonas y reordenar paradas a mano.",
     "Reordenar paradas ya existe; el resto entra al módulo de planificación o queda como próximo paso."),
]:
    row = src.add_row().cells
    cell_text(row[0], fuente, 9.5, True)
    cell_text(row[1], aprendimos, 9)
    cell_text(row[2], decidimos, 9)
    shade(row[2], "F1F8D9")
fix_widths(src, w)

h1(doc, "4", "Cómo cambió la idea")
para(doc, "Cada fila indica su estado real al cierre de este informe.", 9.5, False, GREY, after=6)
ch = doc.add_table(rows=1, cols=4)
borders(ch, "E2E8F0", "4")
w = (Cm(3.0), Cm(5.0), Cm(6.0), Cm(2.6))
header_row(ch, ("Tema", "Al llegar", "Ahora", "Estado"), w)
HECHO, CURSO, PROX = ("HECHO", LIME), ("EN CURSO", CYAN), ("PRÓXIMO", ORANGE)
for tema, antes, ahora, estado in [
    ("Prioridad", "Riesgo por cobertura de días", "Además, hora límite para puntos críticos y pedidos ordenados por prioridad y ventana del punto", HECHO),
    ("Capacidad", "La ruta asumía que todo cabía", "Carga parcial según capacidad; lo que no cabe queda como pendiente con su cantidad", HECHO),
    ("Agente", "Explicaba decisiones", "Consulta stock por bodega, origen y vehículos disponibles, y entrega una propuesta única con costo y su incertidumbre", HECHO),
    ("Datos", "Esperábamos datos del evento", "Revisión automática de viabilidad de datasets y dataset propio de Quito, todo sintético", CURSO),
    ("Campo", "Modo sin conexión con fallas", "Bodega y retiro del conductor funcionan sin conexión", HECHO),
    ("Foco del producto", "Plataforma completa", "Módulo de planificación: compara flota propia y contratada y recomienda con justificación", CURSO),
    ("Costo", "Kilómetros × tarifa + horas", "Costo total de servir: combustible según carga, operarios de descarga, merma y precio de combustible fechado", CURSO),
    ("Medición", "Simulación de 90 días", "Seis indicadores de la charla en el panel, en lenguaje de toda la empresa", CURSO),
    ("Mermas", "Rechazo por daño", "Error de rotulado, mermas por motivo y costo estimado de reentrega", CURSO),
    ("Ciudad", "Ventanas horarias", "Horas pico, restricción horaria para camiones y manifestaciones con retraso parcial", PROX),
    ("Pronóstico", "Media móvil sin medir", "Medir su error y planear estacionalidad con más historia", PROX),
]:
    row = ch.add_row().cells
    cell_text(row[0], tema, 9.5, True)
    cell_text(row[1], antes, 9)
    cell_text(row[2], ahora, 9)
    chip(row[3], estado[0], estado[1], 8)
fix_widths(ch, w)

h1(doc, "5", "El nuevo foco: planificación que decide con números")
callout(doc, [
    ("«Para cada día de entregas, el sistema compara rutas con flota propia y contratada, suma todo lo que cuesta servir "
     "—kilómetros, combustible, personas que descargan y mermas—, revisa horarios, tráfico y restricciones de Quito, "
     "y recomienda una opción con su justificación. La IA explica; el planificador decide.»", False)],
    border=LIME, fill="F1F8D9")
for titulo, texto in [
    ("Costo total de servir: ", "combustible que sube con la carga, desgaste, conductor, ayudantes según el peso y los bultos de cada parada, y merma esperada."),
    ("Propia o contratada: ", "la misma demanda se cotiza con nuestra flota y con un operador externo; si la flota no alcanza, se ve cuánto cuesta cubrir el faltante."),
    ("Recomendación auditable: ", "un puntaje con pesos visibles (costo, críticos a tiempo, pendientes y riesgo) y una justificación con cifras."),
    ("Datos propios: ", "un escenario sintético de Quito con ventanas de centros comerciales, horas pico, manifestaciones y tres tipos de vehículo."),
]:
    bullet(doc, [(titulo, True), (texto, False)], 10)

h1(doc, "6", "Riesgos y estrategias de mitigación")
para(doc, "La charla de logística dio especial peso a este punto. Resumen de los que guían nuestro trabajo:", after=6)
rk = doc.add_table(rows=1, cols=2)
borders(rk, "E2E8F0", "4")
w = (Cm(5.4), Cm(11.2))
header_row(rk, ("Riesgo", "Mitigación"), w)
for riesgo, mitigacion in [
    ("Datos incompletos o inconsistentes", "Validación al importar, marca «sin dato» en vez de inventar, catálogo estandarizado y conteos físicos por muestreo."),
    ("Resistencia al cambio", "Pantallas simples por rol, trabajo sin conexión, capacitación corta por rol y control humano de cada decisión."),
    ("Rutas poco realistas", "Validar capacidad, ventanas, jornada y restricciones antes de publicar; mostrar el motivo de cada descarte."),
    ("Pronósticos imprecisos", "Modelo de referencia simple y explicable; medir su error y ajustar parámetros periódicamente."),
    ("Tráfico, cierres y manifestaciones", "Incidentes con radio de impacto y replanificación con aprobación: llegar tarde o atender primero a los críticos."),
    ("Flota propia insuficiente", "Opción con operador contratado y su costo comparado."),
    ("La IA se equivoca", "La IA solo explica; el cálculo es determinista y existe una explicación de respaldo."),
    ("Costos de implementación", "MVP enfocado en planificación, con indicadores medibles desde el primer día."),
]:
    row = rk.add_row().cells
    cell_text(row[0], riesgo, 9.5, True)
    cell_text(row[1], mitigacion, 9)
fix_widths(rk, w)

h1(doc, "7", "Cómo trabajamos hoy")
for titulo, texto in [
    ("Escucha organizada: ", "cada integrante registró charlas y preguntas en notas colaborativas con hora, fuente y nivel de certeza."),
    ("Del aprendizaje al backlog: ", "las ideas del día se analizaron contra el código existente y se convirtieron en issues con prioridad (alta, media, baja) y responsable."),
    ("Coordinación: ", "un tablero fijado define la cola de cada persona y el orden en que se integran los archivos compartidos, para no pisarnos."),
    ("Calidad: ", "cada cambio entra por pull request revisado y con pruebas; la documentación se actualiza en la misma bóveda."),
    ("Roles: ", "Bryan (motor, datos y agente), Anthony (motor de rutas y mermas), Alexander (planificador web y campo), "
                "Jorge (infraestructura, despliegue y panel) y Joseph (coordinación, documentación y pitch)."),
]:
    bullet(doc, [(titulo, True), (texto, False)], 10)

h1(doc, "8", "Próximos pasos hasta la entrega")
nx = doc.add_table(rows=1, cols=2)
borders(nx, "E2E8F0", "4")
w = (Cm(4.2), Cm(12.4))
header_row(nx, ("Cuándo", "Qué"), w)
for cuando, que in [
    ("Noche del viernes", "Costo total de servir, operador contratado, comparación A/B/C en pantalla, dataset propio e indicadores."),
    ("Sábado 09:00–10:00", "Ronda de pruebas de punta a punta por roles."),
    ("Sábado 11:30", "Congelamiento del código; lo que no llegue se presenta como próximo paso."),
    ("Sábado 11:30–13:15", "Validación integral con mentores sobre el prototipo."),
    ("Sábado 15:45 y 16:00", "Congelamiento de materiales y entrega de hitos."),
]:
    row = nx.add_row().cells
    cell_text(row[0], cuando, 9.5, True)
    cell_text(row[1], que, 9)
fix_widths(nx, w)

para(doc, "", after=6)
p = para(doc, "", after=2)
rich(p, [("Declaración de veracidad. ", True),
         ("Los cambios descritos como hechos están integrados en nuestro repositorio; "
          "los marcados «en curso» y «próximo» todavía no. Los escenarios de datos son sintéticos y no hay cifras de ahorro confirmadas por una empresa.", False)], 9.5)
left_bar(p, ORANGE, "24")

out = AQUI / "IstpetDev-Informe-Checkpoint-16oct.docx"
doc.save(out)
print(out)
