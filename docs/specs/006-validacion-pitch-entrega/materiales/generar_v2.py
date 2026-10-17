"""Genera pitch-v2.pptx y pitch-v2.pdf desde afirmaciones T4 verificadas.

Uso:
    uv run --with python-pptx python docs/specs/006-validacion-pitch-entrega/materiales/generar_v2.py

LibreOffice debe estar disponible como `soffice` para exportar el PDF.
"""
from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
ASSETS = HERE / "entregable-2" / "img"
ROUTE_CAPTURE = ASSETS / "captura-rutas.png"
PPTX_PATH = HERE / "pitch-v2.pptx"
PDF_PATH = HERE / "pitch-v2.pdf"

W, H = 13.333, 7.5
CYAN = RGBColor(0x00, 0xA7, 0xC4)
LIME = RGBColor(0x9E, 0xC2, 0x00)
ORANGE = RGBColor(0xEA, 0x60, 0x1F)
INK = RGBColor(0x2D, 0x37, 0x48)
MUTED = RGBColor(0x55, 0x63, 0x73)
PALE = RGBColor(0xF4, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD8, 0xE2, 0xE7)
PALE_CYAN = RGBColor(0xE8, 0xF7, 0xFA)
PALE_LIME = RGBColor(0xF4, 0xF8, 0xE3)
PALE_ORANGE = RGBColor(0xFD, 0xF0, 0xE9)
FONT = "Liberation Sans"


def rect(slide, x, y, w, h, fill, line=LINE, radius=True, line_width=1.0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    return shape


def text(slide, x, y, w, h, value, size=20, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.02,
         font=FONT, auto_fit=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    if auto_fit:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    paras = value if isinstance(value, list) else [value]
    for i, item in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.alignment = align
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.space_after = Pt(5)
    return box


def bullet_list(slide, x, y, w, h, items, size=21, color=INK, gap=14):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
    return box


def new_slide(prs, section, title, source):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = PALE
    rect(slide, 0, 0, W, 0.13, CYAN, CYAN, radius=False)
    text(slide, 0.62, 0.30, 5.6, 0.34, section.upper(), 14, CYAN, True)
    text(slide, 0.62, 0.76, 12.0, 0.68, title, 29, INK, True)
    rect(slide, 0.62, 6.84, 12.08, 0.015, LINE, LINE, radius=False)
    text(slide, 0.66, 6.90, 12.0, 0.50, source, 14, MUTED)
    return slide


def card(slide, x, y, w, h, heading, body, accent=CYAN,
         fill=WHITE, heading_size=20, body_size=18):
    rect(slide, x, y, w, h, fill)
    rect(slide, x, y, 0.08, h, accent, accent, radius=False)
    text(slide, x + 0.22, y + 0.18, w - 0.42, 0.42, heading,
         heading_size, INK, True)
    text(slide, x + 0.22, y + 0.70, w - 0.42, h - 0.84, body,
         body_size, MUTED)


def slide_problem(prs):
    s = new_slide(prs, "Anexo 5.7 · Problema", "¿Flota propia o apoyo contratado?",
                  "Fuente: guion T4, hipótesis de problema A-02; decisión de producto P-02.")
    text(s, 0.82, 1.70, 11.6, 0.68,
         "Cada día hay que decidir cómo servir las entregas disponibles.",
         25, INK, True, align=PP_ALIGN.CENTER)
    card(s, 0.86, 2.70, 5.55, 2.25, "Mirar solo kilómetros", 
         "puede ocultar combustible, trabajo de descarga y producto expuesto a merma.",
         ORANGE, PALE_ORANGE, 22, 21)
    card(s, 6.90, 2.70, 5.55, 2.25, "Comparar el servicio completo",
         "ayuda a que el planificador entienda los costos y las entregas pendientes.",
         CYAN, PALE_CYAN, 22, 21)
    text(s, 1.12, 5.48, 11.05, 0.72,
         "Nuestra hipótesis: una decisión explicable entre flota propia y contratada.",
         21, INK, True, align=PP_ALIGN.CENTER)


def slide_solution(prs):
    s = new_slide(prs, "Anexo 5.7 · Solución", "El motor calcula. La persona decide.",
                  "Fuente: matriz A-10, A-11 y P-01/P-02; guion T4.")
    labels = [
        ("Comparar", "Rutas con flota propia y contratada", CYAN, PALE_CYAN),
        ("Desglosar", "Costo total estimado de servir", LIME, PALE_LIME),
        ("Explicar", "La IA comunica los datos calculados", ORANGE, PALE_ORANGE),
        ("Decidir", "El planificador elige la opción", CYAN, WHITE),
    ]
    start_x, gap, cw = 0.78, 0.25, 2.82
    for i, (head, body, accent, fill) in enumerate(labels):
        x = start_x + i * (cw + gap)
        card(s, x, 2.15, cw, 2.25, head, body, accent, fill, 20, 18)
        if i < len(labels) - 1:
            text(s, x + cw, 2.90, gap, 0.45, "›", 27, MUTED, True,
                 align=PP_ALIGN.CENTER)
    text(s, 0.90, 4.55, 11.45, 0.34,
         "PROPUESTA · la comparación completa A/B/C sigue en integración",
         16, ORANGE, True, align=PP_ALIGN.CENTER)
    rect(s, 1.42, 5.08, 10.5, 0.88, PALE_CYAN, CYAN)
    text(s, 1.65, 5.30, 10.05, 0.42,
         "La IA explica; el planificador decide.", 25, INK, True,
         align=PP_ALIGN.CENTER)


def slide_demo_capture(prs):
    s = new_slide(prs, "Anexo 5.7 · Prototipo / demo", "Una captura real de la pantalla de rutas",
                  "Fuente: captura-rutas.png, Entregable 2. La captura muestra otro escenario sintético; no es Quito ni A/B/C.")
    if ROUTE_CAPTURE.exists():
        rect(s, 0.76, 1.72, 7.12, 3.02, WHITE, LINE, radius=False, line_width=1.2)
        # Recorta navegación, métricas y fecha; conserva el mapa como evidencia visual de la app.
        pic = s.shapes.add_picture(str(ROUTE_CAPTURE), Inches(0.78), Inches(1.74),
                                   width=Inches(7.08), height=Inches(2.98))
        pic.crop_left = 0.415
        pic.crop_right = 0.207
        pic.crop_top = 0.397
        pic.crop_bottom = 0.308
    card(s, 8.35, 1.82, 4.10, 2.92, "Interfaz existente",
         "La captura documenta una vista de rutas. No demuestra la comparación de flota ni una corrida del escenario de Quito.",
         CYAN, WHITE, 21, 19)
    rect(s, 0.78, 5.08, 11.68, 0.90, PALE_ORANGE, ORANGE)
    text(s, 1.00, 5.27, 11.24, 0.46,
         "Captura real · los datos visibles pertenecen a otro escenario sintético.",
         19, INK, True, align=PP_ALIGN.CENTER)


def slide_abc(prs):
    s = new_slide(prs, "Anexo 5.7 · Prototipo / demo", "Diseño del panel de comparación A / B / C",
                  "Fuente: propuesta del evento y matriz P-02/P-03/P-05. Celdas sin dato: no atribuir cifras al motor.")
    rect(s, 0.76, 1.58, 11.82, 0.55, PALE_ORANGE, ORANGE)
    text(s, 0.96, 1.70, 11.40, 0.30,
         "DISEÑO · comparación contratada y recomendación calculada siguen en integración",
         16, INK, True, align=PP_ALIGN.CENTER)

    x0, y0 = 0.82, 2.32
    widths = [2.46, 2.52, 2.52, 2.52]
    headings = ["Componente", "A · Flota propia", "B · Propia, críticos primero", "C · Contratada"]
    xx = x0
    for j, (w, title) in enumerate(zip(widths, headings)):
        rect(s, xx, y0, w, 0.66, CYAN if j == 0 else INK,
             CYAN if j == 0 else INK, radius=False)
        text(s, xx + 0.07, y0 + 0.12, w - 0.14, 0.42, title, 15,
             WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        xx += w

    rows = ["Combustible", "Desgaste", "Conductor y ayudantes", "Merma esperada", "Costo total"]
    rh = 0.48
    for i, label in enumerate(rows):
        yy = y0 + 0.66 + i * rh
        fill = WHITE if i % 2 == 0 else PALE_CYAN
        xx = x0
        for j, w in enumerate(widths):
            rect(s, xx, yy, w, rh, fill, LINE, radius=False, line_width=0.6)
            value = label if j == 0 else "sin dato"
            text(s, xx + 0.07, yy + 0.10, w - 0.14, 0.27, value,
                 15 if j else 16, INK if j == 0 else MUTED,
                 bold=(j == 0 or i == len(rows)), align=PP_ALIGN.CENTER,
                 valign=MSO_ANCHOR.MIDDLE)
            xx += w

    rect(s, 0.82, 5.50, 7.48, 0.66, PALE_LIME, LIME)
    text(s, 1.00, 5.59, 7.12, 0.46,
         "Recomendada: sin dato · el planificador decide",
         17, INK, True, valign=MSO_ANCHOR.MIDDLE)
    rect(s, 8.52, 5.50, 4.12, 0.66, PALE, LINE)
    text(s, 8.68, 5.57, 3.80, 0.50,
         "Justificación: pendiente de integrar el puntaje A/B/C.",
         16, INK, False, valign=MSO_ANCHOR.MIDDLE)
    rect(s, 0.82, 6.25, 11.82, 0.48, PALE_ORANGE, ORANGE)
    text(s, 1.00, 6.30, 11.45, 0.34,
         "[CAPTURA PENDIENTE: pantalla A/B/C]",
         17, INK, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def slide_impact_indicators(prs):
    s = new_slide(prs, "Anexo 5.7 · Impacto", "Indicadores que cualquier rol puede entender",
                  "Propuesta de indicadores: matriz P-07 en curso; diseño de medición según la propuesta de planificación.")
    text(s, 0.82, 1.58, 11.6, 0.44,
         "Preguntas sencillas para conversar con el equipo operativo",
         20, MUTED, False)
    items = [
        ("Costo por kilómetro", "¿Cuánto cuesta recorrer esta ruta?", CYAN),
        ("Costo por entrega", "¿Cuánto cuesta atender cada pedido?", LIME),
        ("Entregas a tiempo", "¿Llegamos dentro de la ventana acordada?", ORANGE),
        ("Pendientes", "¿Qué pedidos quedaron sin atender?", CYAN),
        ("Mermas", "¿Qué producto se dañó o rechazó?", LIME),
        ("Propia frente a contratada", "¿Qué opción conviene con los datos disponibles?", ORANGE),
    ]
    coords = [(0.82, 2.18), (4.88, 2.18), (8.94, 2.18),
              (0.82, 4.24), (4.88, 4.24), (8.94, 4.24)]
    for (head, body, accent), (x, y) in zip(items, coords):
        card(s, x, y, 3.58, 1.68, head, body, accent, WHITE, 17, 16)


def slide_viability(prs):
    s = new_slide(prs, "Anexo 5.7 · Viabilidad", "Qué podemos respaldar hoy",
                  "Fuente: matriz P-01/P-04/P-11 y A-22. Ensayo P-01: sintético, fixture demo-base; no es la demo A/B/C.")
    card(s, 0.82, 1.86, 5.65, 2.02, "Costo estimado de servir",
         "Ensayo técnico sintético: combustible según carga, desgaste, conductor, ayudantes y merma esperada.",
         CYAN, WHITE, 21, 19)
    card(s, 6.86, 1.86, 5.65, 2.02, "Importación y escenario",
         "Hay prueba de importación del escenario sintético de Quito. La comparación completa no está ensayada en pantalla.",
         LIME, WHITE, 21, 19)
    rect(s, 0.82, 4.40, 11.69, 1.12, PALE_ORANGE, ORANGE)
    text(s, 1.08, 4.61, 11.15, 0.70,
         "Límite: costo estimado y pruebas del motor no equivalen a resultados medidos de una empresa.",
         21, INK, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def slide_business(prs):
    s = new_slide(prs, "Anexo 5.7 · Modelo de negocio", "Una propuesta comercial por validar",
                  "Fuente: matriz A-26/A-30/A-31. Sin precio, ventas ni costos mensuales confirmados.")
    card(s, 0.92, 2.00, 5.46, 2.45, "Cliente propuesto",
         "Empresas de servicios de limpieza con operaciones de abastecimiento y rutas por planificar.",
         CYAN, PALE_CYAN, 22, 20)
    card(s, 6.92, 2.00, 5.46, 2.45, "Cómo se propone cobrar",
         "Implementación e integración inicial, más una suscripción. Precio y disposición de pago por validar.",
         ORANGE, PALE_ORANGE, 22, 20)
    text(s, 1.04, 5.17, 11.2, 0.68,
         "El piloto serviría para validar el uso, los costos reales y el valor para el cliente.",
         21, INK, True, align=PP_ALIGN.CENTER)


def slide_risks(prs):
    s = new_slide(prs, "Anexo 5.7 · Riesgos y mitigación",
                  "RIESGOS Y ESTRATEGIAS DE MITIGACIÓN",
                  "Fuente: matriz P-16 y diseño de riesgos de la propuesta del evento. Medidas propuestas, no todas implementadas.")
    rows = [
        ("Datos incompletos", "Validar al importar; mostrar ‘sin dato’ y supuestos.", CYAN),
        ("Ruta poco realista", "Revisar ventanas, capacidad y restricciones antes de publicar.", ORANGE),
        ("Costo estimado inexacto", "Mostrar el desglose y calibrar con viajes conciliados.", LIME),
        ("Error en la explicación de IA", "Basarla en cifras calculadas; el planificador decide.", CYAN),
        ("Resistencia al cambio", "Capacitar por rol y revisar el plan con el equipo.", ORANGE),
    ]
    y = 1.74
    for risk, mitigation, accent in rows:
        rect(s, 0.84, y, 3.78, 0.83, WHITE, LINE)
        rect(s, 0.84, y, 0.08, 0.83, accent, accent, radius=False)
        text(s, 1.05, y + 0.16, 3.38, 0.48, risk, 18, INK, True,
             valign=MSO_ANCHOR.MIDDLE)
        text(s, 4.72, y + 0.16, 0.40, 0.48, "→", 20, MUTED, True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        rect(s, 5.24, y, 7.24, 0.83, WHITE, LINE)
        text(s, 5.48, y + 0.13, 6.76, 0.58, mitigation, 17, MUTED,
             valign=MSO_ANCHOR.MIDDLE)
        y += 0.93
    text(s, 0.90, 6.40, 11.5, 0.26,
         "La charla de logística destacó especialmente este análisis.",
         14, MUTED, False)


def slide_roadmap(prs):
    s = new_slide(prs, "Anexo 5.7 · Próximo paso", "Hoja de ruta: validar antes de prometer",
                  "Fuente: matriz P-02/P-05/P-10/P-12/P-13/P-14/P-15 y A-32. Propuestas sujetas a validación.")
    steps = [
        ("Integrar", "Comparación propia / contratada y recomendación justificada", CYAN),
        ("Validar Quito", "Tráfico y restricciones con fuentes vigentes", ORANGE),
        ("Completar el ciclo", "Guía de retorno para entregas pendientes", LIME),
        ("Calibrar", "Pronóstico y costos con datos conciliados", CYAN),
    ]
    for i, (head, body, accent) in enumerate(steps):
        x = 0.82 + (i % 2) * 6.03
        y = 1.82 + (i // 2) * 1.62
        card(s, x, y, 5.55, 1.32, head, body, accent,
             WHITE, 19, 17)
    rect(s, 0.82, 5.28, 11.58, 0.92, PALE_CYAN, CYAN)
    text(s, 1.04, 5.49, 11.14, 0.48,
         "Próximo paso: acordar un piloto, una línea base y los indicadores con una empresa.",
         19, INK, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def build_deck():
    if not ROUTE_CAPTURE.exists():
        raise FileNotFoundError(f"Falta la captura real requerida: {ROUTE_CAPTURE}")
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "IstpetDev · Pitch Expo Clean · v2"
    prs.core_properties.subject = "Planificación logística, basado en guion T4"
    for builder in (
        slide_problem,
        slide_solution,
        slide_demo_capture,
        slide_abc,
        slide_impact_indicators,
        slide_viability,
        slide_business,
        slide_risks,
        slide_roadmap,
    ):
        builder(prs)
    prs.save(PPTX_PATH)

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("No se encontró soffice/libreoffice para generar el PDF")
    if PDF_PATH.exists():
        PDF_PATH.unlink()
    profile = Path("/tmp/lo-profile-pitch-v2")
    profile.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless",
         "--convert-to", "pdf", "--outdir", str(HERE), str(PPTX_PATH)],
        check=True, capture_output=True, text=True,
    )
    if not PDF_PATH.exists():
        raise RuntimeError(f"LibreOffice no produjo el PDF: {result.stdout} {result.stderr}")
    print(f"PPTX: {PPTX_PATH}")
    print(f"PDF:  {PDF_PATH}")


if __name__ == "__main__":
    build_deck()
